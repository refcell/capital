#!/usr/bin/env python3
"""Build the Box, Inc. six-sheet valuation workbook."""
from math import isclose
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUT = Path(__file__).with_name("2026-10-30 Box Model.xlsx")
PRICE = 43.34  # October 29, 2026 close
SHARES = 137.20  # common shares outstanding, millions
PREF_SHARES = 18.566  # Series A preferred as-converted at $27.00, millions
MARKET_CAP = PRICE * SHARES  # millions
CASH = 444.14  # cash and short-term investments, millions
NOTES = 460.0  # 2029 convertible notes principal, millions
LEASES = 102.94  # current + non-current lease liabilities, millions
TOTAL_DEBT = 555.14  # StockAnalysis standardized debt (notes carrying value + leases)
PREF_LIQ = 500.0  # Series A liquidation preference, millions
ENTERPRISE_VALUE = MARKET_CAP + TOTAL_DEBT - CASH
EV_WITH_PREF = ENTERPRISE_VALUE + PREF_SHARES * PRICE
FD_SHARES = SHARES + PREF_SHARES
FD_MARKET_CAP = FD_SHARES * PRICE
TTM_REVENUE = 1234.0
TTM_FCF_SA = 389.17  # OCF less capex
TTM_FCF_NONGAAP = 346.3  # company definition, also nets capitalized software
TTM_SBC = 235.74
FY27_REVENUE = 1290.0  # management guide and consensus
FY27_EPS = 1.54  # non-GAAP diluted EPS guide and consensus
FY27_FCF = 362.87  # StockAnalysis consensus

# Five-year discounted terminal FCF framework on company-defined (non-GAAP) FCF.
# Shares are fully diluted for the in-the-money Series A preferred (18.6M shares);
# the 2029 notes are treated as debt because the capped call offsets dilution to $66.86.
SCENARIOS = {
    "Bear": {"rev_cagr": 0.06, "fcf_margin": 0.27, "multiple": 12.0, "net_cash": -250.0, "shares": 152.0, "discount": 0.100, "weight": 0.25},
    "Base": {"rev_cagr": 0.10, "fcf_margin": 0.30, "multiple": 15.0, "net_cash": -100.0, "shares": 142.0, "discount": 0.090, "weight": 0.50},
    "Bull": {"rev_cagr": 0.13, "fcf_margin": 0.33, "multiple": 18.0, "net_cash": 0.0, "shares": 136.0, "discount": 0.085, "weight": 0.25},
}
for case in SCENARIOS.values():
    case["terminal_revenue"] = FY27_REVENUE * (1 + case["rev_cagr"]) ** 5
    case["terminal_fcf"] = case["terminal_revenue"] * case["fcf_margin"]
    case["terminal_ev"] = case["terminal_fcf"] * case["multiple"]
    case["terminal_price"] = (case["terminal_ev"] + case["net_cash"]) / case["shares"]
    case["target"] = case["terminal_price"] / (1 + case["discount"]) ** 5
weighted_value = sum(case["target"] * case["weight"] for case in SCENARIOS.values())

ANALYST_AVG_PT = 39.86
assert SCENARIOS["Bear"]["target"] < PRICE * 0.9
assert abs(SCENARIOS["Base"]["target"] - ANALYST_AVG_PT) / ANALYST_AVG_PT < 0.20

rf = 5.05
erp = 5.00
beta = 0.72
cost_equity = rf + beta * erp
pretax_debt_cost = 5.50
tax_rate = 21.0
equity_weight = MARKET_CAP / (MARKET_CAP + TOTAL_DEBT)
debt_weight = 1 - equity_weight
wacc = equity_weight * cost_equity + debt_weight * pretax_debt_cost * (1 - tax_rate / 100)

wb = Workbook()
navy, blue, gold, gray, white = "17365D", "D9EAF7", "D8B34B", "E7E6E6", "FFFFFF"
thin = Side(style="thin", color="A6A6A6")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def put(ws, row, col, value, *, bold=False, fill=None, color="000000", size=10, wrap=True):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(name="Aptos", size=size, bold=bold, color=color)
    cell.fill = PatternFill("solid", fgColor=fill) if fill else PatternFill(fill_type=None)
    cell.border = border
    cell.alignment = Alignment(vertical="top", wrap_text=wrap)
    return cell


def title(ws, text, end_col=5, subtitle=None):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    put(ws, 1, 1, text, bold=True, fill=navy, color=white, size=15)
    ws.row_dimensions[1].height = 28
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
        put(ws, 2, 1, subtitle, bold=True, fill=blue)


def header(ws, row, values):
    for col, value in enumerate(values, 1):
        put(ws, row, col, value, bold=True, fill=navy, color=white)


def widths(ws, values):
    for col, width in enumerate(values, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A3"


# 1. Valuation
ws = wb.active
ws.title = "Valuation"
title(ws, "Box, Inc. (BOX) — Valuation", 4, "Quote date: October 29, 2026 close | Model date: October 30, 2026")
facts = [
    ("Company", "Box, Inc.", "Intelligent Content Management (ICM) SaaS: cloud content management, security, e-signature, workflow and Box AI"),
    ("Ticker", "NYSE: BOX", "Technology / Software - Infrastructure; fiscal year ends January 31"),
    ("Price", PRICE, "StockAnalysis / Fiscal.ai October 29, 2026 close; 52-week range $21.34-$43.57"),
    ("Common shares outstanding (M)", SHARES, "StockAnalysis statistics; down 4.06% YoY from repurchases"),
    ("Series A preferred as-converted (M)", PREF_SHARES, "KKR $500M preferred convertible at $27.00; 18,565,810 shares per May 2026 proxy"),
    ("Fully diluted shares (M)", FD_SHARES, "Common plus preferred as-converted; 2029 notes treated as debt (capped call to $66.86)"),
    ("Market capitalization ($M)", MARKET_CAP, "Common only; StockAnalysis shows $5.95B"),
    ("Fully diluted market cap ($M)", FD_MARKET_CAP, "Includes preferred at as-converted value"),
    ("Enterprise value ($M)", ENTERPRISE_VALUE, "Market cap + $555.1M debt and leases - $444.1M cash; matches StockAnalysis $6.06B"),
    ("EV incl. preferred as-converted ($M)", EV_WITH_PREF, "Preferred is in the money ($43.34 vs. $27.00 conversion price), so it is an equity claim worth ~$805M"),
    ("Primary valuation lens", "Discounted terminal FCF", "GAAP net income is distorted by deferred-tax valuation allowance releases; FCF and non-GAAP EPS are cross-checked"),
    ("Stance", "Watch / Hold", "Fundamentals are accelerating, but a 37% 30-day rerating has pulled the price above the $39.86 average target and to base-case value"),
]
header(ws, 3, ["Field", "Value", "Comment"])
for row, values in enumerate(facts, 4):
    for col, value in enumerate(values, 1):
        cell = put(ws, row, col, value, bold=(col == 1))
        if col == 2 and isinstance(value, float):
            cell.number_format = "$0.00" if values[0] == "Price" else "#,##0.0"

metric_start = 4 + len(facts) + 1
header(ws, metric_start, ["Valuation metric", "Value", "Interpretation"])
metrics = [
    ("Trailing P/E (GAAP)", 62.59, "Misleading: TTM GAAP EPS $0.69 after preferred allocation; FY25/FY26 included tax-allowance releases"),
    ("Forward P/E (FY27 non-GAAP $1.54)", PRICE / FY27_EPS, "Management guide equals consensus; excludes SBC (~18.5% of revenue)"),
    ("P/S (TTM)", MARKET_CAP / TTM_REVENUE, "$1.234B TTM revenue; growth reaccelerated to 9-11%"),
    ("P/FCF (StockAnalysis FCF)", MARKET_CAP / TTM_FCF_SA, "$389.2M OCF less capex; excludes capitalized software"),
    ("P/FCF (company non-GAAP FCF)", MARKET_CAP / TTM_FCF_NONGAAP, "$346.3M TTM after capitalized software"),
    ("EV/FCF (non-GAAP)", ENTERPRISE_VALUE / TTM_FCF_NONGAAP, "Primary current cash-flow cross-check"),
    ("EV/Sales", ENTERPRISE_VALUE / TTM_REVENUE, "Mid-single-digit sales multiple for ~80% gross margin software"),
    ("EV/EBITDA (GAAP)", 45.95, "GAAP EBITDA ($131.8M) is burdened by SBC; weak denominator"),
    ("FCF yield (non-GAAP, common cap)", TTM_FCF_NONGAAP / MARKET_CAP, "Headline cash yield before SBC"),
    ("SBC-adjusted FCF yield (FD cap)", (TTM_FCF_NONGAAP - TTM_SBC) / FD_MARKET_CAP, "Owner-cash view: $346.3M FCF less $235.7M SBC over fully diluted cap"),
    ("Interest coverage", 12.40, "StockAnalysis; 1.50% coupon converts keep cash interest small"),
]
for row, values in enumerate(metrics, metric_start + 1):
    for col, value in enumerate(values, 1):
        cell = put(ws, row, col, value, bold=(col == 1))
        if col == 2 and "yield" in values[0]:
            cell.number_format = "0.0%"
        elif col == 2:
            cell.number_format = "0.0x" if values[0] != "Interest coverage" else "0.0x"
widths(ws, [36, 24, 96, 14])

# 2. WACC
ws = wb.create_sheet("WACC")
title(ws, "Box — WACC", 4, "CAPM and debt-weighted cost of capital")
header(ws, 3, ["Component", "Value", "Source / formula"])
wacc_rows = [
    ("Risk-free rate", rf / 100, "FRED DGS10 10-year Treasury, October 28, 2026 (5.05%)"),
    ("Equity risk premium", erp / 100, "Model assumption"),
    ("Levered beta", beta, "StockAnalysis five-year beta"),
    ("Cost of equity", cost_equity / 100, "Risk-free rate + beta × ERP"),
    ("Pre-tax cost of debt", pretax_debt_cost / 100, "Economic cost; 2029 notes carry a 1.50% coupon but embed a conversion option"),
    ("Tax rate", tax_rate / 100, "U.S. statutory; cash taxes are low due to NOLs"),
    ("Market cap ($M)", MARKET_CAP, "Price × common shares"),
    ("Total debt ($M)", TOTAL_DEBT, "StockAnalysis: $452.2M notes carrying value + $102.9M leases"),
    ("Equity weight", equity_weight, "Market cap / (market cap + debt)"),
    ("Debt weight", debt_weight, "Debt / (market cap + debt)"),
    ("After-tax debt cost", pretax_debt_cost / 100 * (1 - tax_rate / 100), "Kd × (1 − tax rate)"),
    ("WACC", wacc / 100, "E/V × Ke + D/V × Kd × (1 − t); preferred excluded (3% coupon, in-the-money convert behaves as equity)"),
    ("Scenario discount rates", "10.0% / 9.0% / 8.5%", "Bear / base / bull; above WACC to reflect execution and AI-disruption risk"),
]
for row, values in enumerate(wacc_rows, 4):
    for col, value in enumerate(values, 1):
        cell = put(ws, row, col, value, bold=(col == 1 or values[0] == "WACC"), fill=(gold if values[0] == "WACC" else None))
        if col == 2 and isinstance(value, float) and values[0] not in {"Levered beta", "Market cap ($M)", "Total debt ($M)"}:
            cell.number_format = "0.00%"
        elif col == 2 and values[0] in {"Market cap ($M)", "Total debt ($M)"}:
            cell.number_format = "#,##0.0"
widths(ws, [30, 22, 96, 14])

# 3. Scenarios
ws = wb.create_sheet("Scenarios")
title(ws, "Box — Discounted Terminal FCF Scenarios", 5, "FY2027 (Jan-2027) guide compounded five years to FY2032; terminal value discounted five years; USD millions")
header(ws, 3, ["Metric", "Bear", "Base", "Bull", "Notes"])
keys = ("Bear", "Base", "Bull")
scenario_rows = [
    ("Revenue CAGR (5Y)", *(SCENARIOS[k]["rev_cagr"] for k in keys), "FY2027 base is the $1.290B guide; Q2 grew 9% reported and 11% constant currency"),
    ("Terminal revenue ($M)", *(SCENARIOS[k]["terminal_revenue"] for k in keys), "FY2032 revenue"),
    ("Adjusted FCF margin", *(SCENARIOS[k]["fcf_margin"] for k in keys), "Company non-GAAP FCF; FY27 consensus is ~28.1%; bear assumes AI compute and pricing pressure"),
    ("Terminal FCF ($M)", *(SCENARIOS[k]["terminal_fcf"] for k in keys), "Terminal revenue × FCF margin"),
    ("Exit FCF multiple", *(SCENARIOS[k]["multiple"] for k in keys), "12x / 15x / 18x on SBC-inclusive FCF (SBC ~18% of revenue is an owner cost)"),
    ("Implied EV ($M)", *(SCENARIOS[k]["terminal_ev"] for k in keys), "Terminal FCF × exit multiple"),
    ("Net cash adjustment ($M)", *(SCENARIOS[k]["net_cash"] for k in keys), "Current cash $444M less $460M notes ≈ -$16M; bear/base assume buybacks run ahead of FCF"),
    ("Fully diluted shares (M)", *(SCENARIOS[k]["shares"] for k in keys), "Today 155.8M incl. 18.6M preferred as-converted; repurchases offset SBC at different rates"),
    ("Terminal price", *(SCENARIOS[k]["terminal_price"] for k in keys), "Undiscounted FY2032 value per share"),
    ("Discount rate", *(SCENARIOS[k]["discount"] for k in keys), "Scenario risk adjustment around 8.3% CAPM WACC"),
    ("Present target price", *(SCENARIOS[k]["target"] for k in keys), "Terminal price discounted five years"),
    ("Upside / (downside)", *((SCENARIOS[k]["target"] / PRICE - 1) for k in keys), "Versus $43.34 close"),
    ("Probability", *(SCENARIOS[k]["weight"] for k in keys), "25% / 50% / 25%"),
    ("Weighted value/share", *(SCENARIOS[k]["target"] * SCENARIOS[k]["weight"] for k in keys), "Contribution to fair value"),
]
pct_rows = {"Revenue CAGR (5Y)", "Adjusted FCF margin", "Discount rate", "Upside / (downside)", "Probability"}
usd_rows = {"Terminal price", "Present target price", "Weighted value/share"}
for row, values in enumerate(scenario_rows, 4):
    for col, value in enumerate(values, 1):
        cell = put(ws, row, col, value, bold=(col == 1))
        if col in (2, 3, 4) and values[0] in pct_rows:
            cell.number_format = "0.0%"
        elif col in (2, 3, 4) and values[0] in usd_rows:
            cell.number_format = "$0.00"
        elif col in (2, 3, 4):
            cell.number_format = "#,##0.0"
fv_row = 4 + len(scenario_rows) + 1
put(ws, fv_row, 1, "Probability-weighted fair value", bold=True, fill=gold)
put(ws, fv_row, 2, weighted_value, bold=True, fill=gold).number_format = "$0.00"
put(ws, fv_row + 1, 1, "Upside from current price", bold=True, fill=gold)
put(ws, fv_row + 1, 2, weighted_value / PRICE - 1, bold=True, fill=gold).number_format = "0.0%"
put(ws, fv_row + 2, 1, "Forward P/E cross-check", bold=True, fill=gray)
put(ws, fv_row + 2, 2, "At $43.34 BOX trades at 28.1x FY27 non-GAAP EPS of $1.54. Base target of ~$42 equals ~24x FY28 EPS if EPS compounds ~13% (StockAnalysis 3Y EPS growth forecast 12.8%) to ~$1.74. Analyst targets: $26 low / $39.86 average / $50 high (9 analysts).", fill=gray)
ws.merge_cells(start_row=fv_row + 2, start_column=2, end_row=fv_row + 2, end_column=5)
put(ws, fv_row + 3, 1, "Framework note", bold=True, fill=gray)
put(ws, fv_row + 3, 2, "FCF is used rather than GAAP EPS because FY25 and FY26 GAAP net income include large deferred-tax valuation allowance releases. Company FCF is not owner cash: TTM SBC of $235.7M equals 68% of $346.3M non-GAAP FCF, and TTM repurchases of $460.7M exceeded FCF. Multiples are set below typical SaaS FCF multiples for that reason. Preferred is counted as 18.6M shares because it is deep in the money.", fill=gray)
ws.merge_cells(start_row=fv_row + 3, start_column=2, end_row=fv_row + 3, end_column=5)
ws.row_dimensions[fv_row + 2].height = 45
ws.row_dimensions[fv_row + 3].height = 60
widths(ws, [32, 18, 18, 18, 92])

# 4. Actuals Source Audit
ws = wb.create_sheet("Actuals Source Audit")
title(ws, "Box — Actuals Source Audit", 5, "Financial statement figures in USD millions unless noted; fiscal year ends January 31")
header(ws, 3, ["Data point", "Value", "Source URL", "Source date", "Notes"])
sa = "https://stockanalysis.com/stocks/box"
ir = "https://www.businesswire.com/news/home/20260824340222/en/Box-Reports-Second-Quarter-Fiscal-2027-Financial-Results"
remarks = "https://s21.q4cdn.com/328470014/files/doc_earnings/2027/q2/transcript/Box-Q2FY27-earnings-prepared-remarks.pdf"
deck = "https://s21.q4cdn.com/328470014/files/doc_earnings/2027/q2/presentation/Earnings-DECK-Q2FY27.pdf"
pref = "https://www.sec.gov/Archives/edgar/data/1372612/000119312526098466/R20.htm"
notes_url = "https://www.sec.gov/Archives/edgar/data/1372612/000119312526368289/R16.htm"
audit = [
    ("Stock price", "$43.34", "https://fiscal.ai/company/NYSE_BOX/", "2026-10-29", "Close; open $42.00; prior close $42.80; record high zone"),
    ("Market cap", "$5.95B", f"{sa}/statistics/", "2026-10-29", "Common shares only"),
    ("Enterprise value", "$6.06B", f"{sa}/statistics/", "2026-10-29", "Excludes Series A preferred"),
    ("Shares outstanding", "137.20M", f"{sa}/statistics/", "2026-10-29", "-4.06% YoY; float 132.4M"),
    ("Revenue TTM", "$1,234.1M", f"{sa}/financials/", "2026-07-31", "+9.63%; FY26 $1,177M (+7.99%)"),
    ("Revenue by geography TTM", "US $771.4M / Japan $314.8M / Other $147.9M", f"{sa}/financials/", "2026-07-31", "Japan is 25.5% of revenue; ~70% of non-US revenue is yen"),
    ("Gross profit TTM", "$981.7M", f"{sa}/financials/", "2026-07-31", "79.55% GAAP margin"),
    ("Operating income TTM", "$125.6M", f"{sa}/financials/", "2026-07-31", "10.17% GAAP margin; FY26 $92.4M, FY25 $79.6M, FY24 $50.8M"),
    ("Net income to common TTM", "$100.5M", f"{sa}/financials/", "2026-07-31", "Total net income $130.7M; preferred allocation deducted"),
    ("GAAP EPS TTM", "$0.69", f"{sa}/financials/", "2026-07-31", "FY25 $1.36 included tax-allowance release"),
    ("Income tax TTM", "-$9.48M (benefit)", f"{sa}/statistics/", "2026-07-31", "Deferred tax asset $268.1M on balance sheet"),
    ("Cash and investments", "$444.1M", f"{sa}/financials/balance-sheet/", "2026-07-31", "$342.9M cash + $101.3M short-term investments"),
    ("Total debt", "$555.1M", f"{sa}/financials/balance-sheet/", "2026-07-31", "$452.2M notes carrying + $102.9M leases"),
    ("2029 convertible notes", "$460.0M, 1.50%, conv. $43.46", notes_url, "2026-07-31", "Capped call strike $43.46, cap $66.86, ~10.6M shares"),
    ("Series A preferred", "$500M liq. pref., 3%, conv. $27.00", pref, "2026-01-31", "18,565,810 as-converted shares per May 2026 proxy; $15M/yr cash dividend"),
    ("Deferred revenue", "$589.5M", ir, "2026-07-31", "Current; down from $647.9M at Jan-26 on seasonality"),
    ("Stockholders' deficit", "-$351.1M", ir, "2026-07-31", "Total equity incl. preferred mezzanine $146.3M"),
    ("Operating cash flow TTM", "$394.5M", f"{sa}/financials/cash-flow-statement/", "2026-07-31", "+16.8%"),
    ("Capex TTM", "$5.3M", f"{sa}/financials/cash-flow-statement/", "2026-07-31", "Plus $39.5M capitalized software / intangibles"),
    ("FCF TTM (StockAnalysis)", "$389.2M", f"{sa}/financials/cash-flow-statement/", "2026-07-31", "31.5% margin; OCF less capex"),
    ("Non-GAAP FCF TTM (company)", "$346.3M", ir, "2026-07-31", "FY26 $312.9M + H1 FY27 $187.5M - H1 FY26 $154.1M"),
    ("SBC TTM", "$235.7M", f"{sa}/financials/cash-flow-statement/", "2026-07-31", "19.1% of revenue; 68% of non-GAAP FCF"),
    ("Repurchases TTM", "$460.7M", f"{sa}/financials/cash-flow-statement/", "2026-07-31", "Exceeds FCF; $378M authorization remaining at Jul-31"),
    ("Q2 FY27 revenue", "$321.1M", ir, "2026-07-31", "+9% reported, +11% constant currency"),
    ("Q2 FY27 RPO / cRPO", "$1.692B / $904.7M", ir, "2026-07-31", "+15% / +11% YoY"),
    ("Q2 FY27 non-GAAP operating margin", "29.4%", ir, "2026-07-31", "GAAP 10.2%"),
    ("Q2 FY27 EPS", "GAAP $0.09 / non-GAAP $0.40", ir, "2026-07-31", "$0.04 FX headwind"),
    ("Net retention rate", "106%", deck, "2026-07-31", "103% Q2 FY26 -> 104% -> 104% -> 105% -> 106%; churn 3%"),
    ("Billings growth Q2", "17%", remarks, "2026-07-31", "Record Q2 bookings"),
    ("FY27 guidance", "Rev ~$1.290B; non-GAAP OM 28%; EPS $1.54 non-GAAP / $0.38 GAAP", ir, "2026-08-25", "141M diluted shares; ~$0.09 FX headwind"),
    ("Q3 FY27 guidance", "Rev ~$329M; EPS $0.39 non-GAAP / $0.12 GAAP", ir, "2026-08-25", "142M diluted shares"),
    ("FY27 consensus", "Rev $1.29B; EPS $1.54; FCF $362.9M", f"{sa}/forecast/", "2026-10-29", "8 analysts; FY28+ gated"),
    ("Analyst price target", "$39.86 avg / $40 median / $26-$50", f"{sa}/forecast/", "2026-10-29", "9 analysts: 3 Strong Buy, 1 Buy, 4 Hold, 1 Sell"),
    ("Beta", "0.72", f"{sa}/statistics/", "2026-10-29", "Five-year"),
    ("Short interest", "17.30M (12.61% of shares)", f"{sa}/statistics/", "2026-10-29", "6.7 days to cover"),
    ("U.S. 10Y Treasury", "5.05%", "https://fred.stlouisfed.org/series/DGS10", "2026-10-28", "Risk-free rate"),
    ("Next earnings", "Dec 1, 2026 (estimated)", f"{sa}/statistics/", "2026-10-29", "Investor product briefing Nov 5, 2026 (no financial update)"),
]
for row, values in enumerate(audit, 4):
    for col, value in enumerate(values, 1):
        put(ws, row, col, value, bold=(col == 1))
widths(ws, [32, 34, 90, 14, 70])

# 5. Questions
ws = wb.create_sheet("Questions")
title(ws, "Box — Open Questions", 4, "Items that can change valuation or conviction")
header(ws, 3, ["#", "Question", "Why it matters", "Best evidence / next check"])
questions = [
    (1, "What are the KKR Series A preferred redemption rights and timing, and will Box force conversion if the VWAP exceeds 200% of the $27.00 conversion price ($54)?", "The preferred is an in-the-money claim on ~18.6M shares (~$805M) and receives $15M/yr in cash dividends; a redemption or conversion changes share count and cash.", "10-K preferred note; Certificate of Designations; any 8-K on conversion."),
    (2, "How will the 2029 notes ($460M, conversion price $43.46) be settled now that the stock is at the conversion price?", "The capped call offsets dilution up to $66.86, but settlement method (cash vs. shares) affects liquidity before September 2029.", "Debt note; Q3 10-Q; management commentary on refinancing."),
    (3, "How much of the 106% net retention rate comes from price-per-seat uplift on Enterprise Advanced versus seat expansion?", "Price-driven NRR is a one-time step-up as the base migrates; seat-driven NRR is more durable.", "Suite mix disclosure; Enterprise Advanced penetration of the base."),
    (4, "How large is AI-unit and platform consumption revenue, and what is its gross margin?", "FY27 gross margin guide (~80.5%, Q4 ~80%) signals AI compute and public-cloud costs; monetization must outrun inference cost.", "Gross margin trajectory; AI units disclosure at BoxWorks briefing Nov 5."),
    (5, "Is 15% RPO growth versus 9-11% revenue growth driven by longer contract durations rather than higher annual value?", "Management attributes RPO strength partly to durations; cRPO grew 11%, a cleaner leading indicator.", "cRPO vs. RPO gap each quarter."),
    (6, "How much of TTM GAAP net income and the $268M deferred tax asset came from valuation allowance releases?", "FY25 Q4 net income of $168M and FY26 Q4 net income of $68.5M were tax-driven; GAAP P/E is unusable.", "Income tax footnote; effective tax rate normalization."),
    (7, "Can SBC fall from ~19% of revenue toward the low teens?", "SBC equals ~68% of non-GAAP FCF; owner FCF is only ~$110M TTM before buybacks.", "SBC guidance (18.5% FY27) and equity burn rate."),
    (8, "Why did repurchases ($460.7M TTM) exceed FCF, and what cash floor will management keep?", "Cash fell from $722.8M to $444.1M in 18 months while notes remain outstanding.", "Buyback pace vs. FCF; $378M remaining authorization."),
    (9, "How exposed is revenue to yen moves, and is Japan growth (+ high-teens constant currency) durable?", "Japan is 25.5% of revenue; FX cost ~100 bps of FY27 growth and $0.09 of EPS.", "Japan revenue in yen; partner (e.g., Japanese resellers) commentary."),
    (10, "Do Microsoft 365 Copilot/SharePoint Premium and Google Workspace bundling erode Box's mid-market seats?", "Bundled AI on incumbent content stores is the main competitive threat to standalone content management.", "Win/loss commentary; churn (currently 3%); seat expansion rate."),
    (11, "What is the revenue contribution and integration status of recent AI tuck-ins and capitalized software ($39.5M TTM)?", "Intangibles rose from $46.8M (FY24) to $103.3M; capitalized development flatters OCF-based FCF.", "Intangible roll-forward; company non-GAAP FCF reconciliation."),
    (12, "Is the Chief Accounting Officer turnover (Berkovitch resigned effective Oct 23; Yuhjtman started Oct 19) purely routine?", "Accounting leadership change near a quarter close warrants attention even though the 8-K states no disagreement.", "Q3 10-Q, auditor comments, any restatements."),
    (13, "Is the 12.6% short interest a hedge against the converts and preferred, or a fundamental short?", "Convertible arbitrage inflates short interest; fundamental shorting would signal skepticism on AI monetization.", "Short interest trend vs. convert hedge ratio."),
    (14, "Can Box reach the 15% revenue growth path cited by BofA, and what is management's long-term margin target?", "The bull case requires mid-teens growth; the base case assumes ~10%.", "BoxWorks investor briefing (Nov 5) and Q3 FY27 results (expected Dec 1)."),
    (15, "When is the next earnings release and will FY27 guidance be raised again?", "Q3 FY27 (expected Dec 1, 2026) tests $329M revenue, 28% margin and NRR guide of 106% exiting FY27.", "Investor-relations calendar and Q3 release."),
]
for row, values in enumerate(questions, 4):
    for col, value in enumerate(values, 1):
        put(ws, row, col, value, bold=(col == 1))
widths(ws, [7, 80, 72, 60])

# 6. Sources
ws = wb.create_sheet("Sources")
title(ws, "Box — Sources", 4, "Public sources accessed October 29-30, 2026")
header(ws, 3, ["#", "Source", "URL", "Use"])
sources = [
    (1, "StockAnalysis overview", f"{sa}/", "Price, range, headline market data"),
    (2, "StockAnalysis financials", f"{sa}/financials/", "Income statement, geography, margins (annual and quarterly)"),
    (3, "StockAnalysis balance sheet", f"{sa}/financials/balance-sheet/", "Cash, investments, debt, leases, deferred revenue, equity"),
    (4, "StockAnalysis cash flow", f"{sa}/financials/cash-flow-statement/", "OCF, capex, FCF, SBC, repurchases, debt repayment"),
    (5, "StockAnalysis statistics", f"{sa}/statistics/", "EV, beta, shares, short interest, valuation ratios"),
    (6, "StockAnalysis forecast", f"{sa}/forecast/", "Consensus revenue/EPS/FCF and price targets"),
    (7, "StockAnalysis profile", f"{sa}/company/", "Company identity, management, recent filings"),
    (8, "Box Q2 FY27 press release (Business Wire)", ir, "Quarterly actuals, guidance, reconciliations"),
    (9, "Box Q2 FY27 prepared remarks", remarks, "Management commentary on NRR, bookings, gross margin and guidance"),
    (10, "Box Q2 FY27 earnings deck", deck, "RPO, NRR trend"),
    (11, "SEC: Box FY26 10-K preferred stock note", pref, "Series A preferred terms"),
    (12, "SEC: Box Q2 FY27 10-Q debt note", notes_url, "2029 convertible notes and capped call terms"),
    (13, "SEC: Box 8-K (CAO resignation), Sept 25, 2026", "https://www.sec.gov/Archives/edgar/data/1372612/000119312526401392/d127648d8k.htm", "Governance"),
    (14, "SEC: Box 8-K (CAO appointment), Oct 6, 2026", "https://www.sec.gov/Archives/edgar/data/1372612/000119312526414757/d176697d8k.htm", "Governance"),
    (15, "Box BoxWorks 2026 investor briefing announcement", "https://www.businesswire.com/news/home/20261016868109/en/Box-to-Host-a-Virtual-Investor-Product-Briefing-at-BoxWorks-2026", "Catalyst timing"),
    (16, "FRED DGS10", "https://fred.stlouisfed.org/series/DGS10", "Risk-free rate"),
    (17, "Fiscal.ai BOX daily prices", "https://fiscal.ai/company/NYSE_BOX/", "October 29 close"),
]
for row, values in enumerate(sources, 4):
    for col, value in enumerate(values, 1):
        put(ws, row, col, value, bold=(col == 1))
widths(ws, [7, 44, 110, 64])

for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = False

wb.save(OUT)
check = load_workbook(OUT, data_only=False)
expected = ["Valuation", "WACC", "Scenarios", "Actuals Source Audit", "Questions", "Sources"]
assert check.sheetnames == expected, check.sheetnames
assert isclose(check["Scenarios"].cell(row=fv_row, column=2).value, weighted_value, rel_tol=1e-12)
print(f"Market cap: ${MARKET_CAP:,.1f}M | EV: ${ENTERPRISE_VALUE:,.1f}M | EV incl. pref: ${EV_WITH_PREF:,.1f}M")
print(f"Cost of equity: {cost_equity:.2f}% | WACC: {wacc:.2f}%")
for name, case in SCENARIOS.items():
    print(f"{name}: terminal ${case['terminal_price']:.2f}, target ${case['target']:.2f} ({case['target'] / PRICE - 1:+.1%})")
print(f"Probability-weighted fair value: ${weighted_value:.2f} ({weighted_value / PRICE - 1:+.1%})")
print(f"Wrote and verified {OUT}")
