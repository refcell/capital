#!/usr/bin/env python3
"""Build the Talen Energy six-sheet valuation workbook."""
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUT = Path(__file__).with_name("2026-09-22 Talen Energy Model.xlsx")
PRICE = 299.97
SHARES = 47.91  # millions
MARKET_CAP = 14.37  # billions
ENTERPRISE_VALUE = 23.71  # billions
DEBT = 9.574  # billions
CASH = 0.232  # billions
NET_DEBT = ENTERPRISE_VALUE - MARKET_CAP
TTM_REVENUE = 3.741  # billions
TTM_EBITDA = 0.582  # GAAP provider figure, billions
TTM_FCF = 0.444  # reported FCF, billions
FY26_REVENUE = 4.53  # consensus, billions
FY26_EPS = 22.26  # adjusted diluted consensus
FY27_REVENUE = 5.24  # visible headline consensus, billions
FY27_EPS = 30.75  # adjusted diluted consensus

# Forward P/E is primary: net debt is 21x reported TTM FCF, while acquisitions
# and adjusted FCF guidance make trailing GAAP cash flow a poor denominator.
SCENARIOS = {
    "Bear": {"rev_cagr": 0.02, "terminal_eps": 22.0, "multiple": 10.0, "weight": 0.20},
    "Base": {"rev_cagr": 0.06, "terminal_eps": 33.0, "multiple": 13.0, "weight": 0.50},
    "Bull": {"rev_cagr": 0.09, "terminal_eps": 42.0, "multiple": 14.0, "weight": 0.30},
}
for case in SCENARIOS.values():
    case["terminal_revenue"] = FY27_REVENUE * (1 + case["rev_cagr"]) ** 5
    case["target"] = case["terminal_eps"] * case["multiple"]
weighted_value = sum(case["target"] * case["weight"] for case in SCENARIOS.values())

assert SCENARIOS["Bear"]["target"] < PRICE
assert abs(SCENARIOS["Base"]["target"] - 459.94) / 459.94 < 0.20
assert 250 < weighted_value < 550

rf = 4.957
erp = 5.00
beta = 1.63
cost_equity = rf + beta * erp
pretax_debt_cost = 6.00
tax_rate = 21.00
equity_weight = MARKET_CAP / (MARKET_CAP + DEBT)
debt_weight = 1 - equity_weight
wacc = equity_weight * cost_equity + debt_weight * pretax_debt_cost * (1 - tax_rate / 100)

wb = Workbook()
navy, blue, gold, gray, white = "17365D", "D9EAF7", "D8B34B", "E7E6E6", "FFFFFF"
thin = Side(style="thin", color="A6A6A6")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def put(ws, row, col, value, *, bold=False, fill=None, color="000000", size=10, wrap=True):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(name="Aptos", size=size, bold=bold, color=color)
    cell.fill = PatternFill("solid", fgColor=fill) if fill else PatternFill(fill_type=None)
    cell.border = border
    cell.alignment = Alignment(vertical="top", wrap_text=wrap)
    return cell


def title(ws, text, end_col=5, subtitle=None):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    put(ws, 1, 1, text, bold=True, fill=navy, color=white, size=15)
    ws.row_dimensions[1].height = 28
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
        put(ws, 2, 1, subtitle, bold=True, fill=blue)


def header(ws, row, values):
    for col, value in enumerate(values, 1):
        put(ws, row, col, value, bold=True, fill=navy, color=white)


def widths(ws, values):
    for col, width in enumerate(values, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A3"


# 1. Valuation
ws = wb.active
ws.title = "Valuation"
title(ws, "Talen Energy (TLN) — Valuation", 4, "Quote date: September 21, 2026 | Model date: September 22, 2026")
header(ws, 3, ["Field", "Value", "Comment"])
facts = [
    ("Company", "Talen Energy Corporation", "Independent power producer with 13.1 GW; PJM-focused nuclear and gas generation"),
    ("Ticker", "NASDAQ: TLN", "Utilities / Independent Power Producers"),
    ("Price", PRICE, "StockAnalysis September 21 close"),
    ("Shares outstanding (M)", SHARES, "StockAnalysis; down 7.07% YoY"),
    ("Market capitalization ($B)", MARKET_CAP, "StockAnalysis"),
    ("Enterprise value ($B)", ENTERPRISE_VALUE, "StockAnalysis"),
    ("Net debt proxy ($B)", NET_DEBT, "Enterprise value less market capitalization"),
    ("Primary valuation lens", "Forward adjusted P/E", "Acquisition-driven scale change and leverage make trailing GAAP P/E and FCF multiples misleading"),
    ("Stance", "Watch / Positive", "Strong forward cash generation and buybacks, balanced by leverage, merchant power exposure and regulatory risk"),
]
for row, values in enumerate(facts, 4):
    for col, value in enumerate(values, 1):
        put(ws, row, col, value, bold=(col == 1))

header(ws, 15, ["Valuation metric", "Value", "Interpretation"])
metrics = [
    ("Trailing P/E", "N/A", "TTM GAAP net loss; not meaningful"),
    ("Forward P/E", 10.05, "Provider forward multiple; primary market lens"),
    ("FY2026 consensus P/E", PRICE / FY26_EPS, "Price divided by non-GAAP adjusted diluted EPS consensus"),
    ("P/S", 3.84, "Acquisition-expanded revenue base makes trailing ratio transitional"),
    ("P/FCF", 32.37, "Reported TTM FCF understates management's adjusted post-acquisition outlook"),
    ("EV/FCF", 53.41, "Shows leverage sensitivity; unsuitable as primary equity framework"),
    ("EV/Sales", 6.34, "High because enterprise value includes $9.34B net debt"),
    ("EV/EBITDA", 40.75, "GAAP provider EBITDA is distorted versus adjusted EBITDA guidance"),
    ("FCF yield", TTM_FCF / MARKET_CAP, "Reported TTM FCF / market cap; adjusted FY2026 guidance implies a much higher yield"),
]
for row, values in enumerate(metrics, 16):
    for col, value in enumerate(values, 1):
        cell = put(ws, row, col, value, bold=(col == 1))
        if values[0] == "FCF yield" and col == 2:
            cell.number_format = "0.0%"
widths(ws, [30, 24, 82, 14])

# 2. WACC
ws = wb.create_sheet("WACC")
title(ws, "Talen Energy — WACC", 4, "CAPM and debt-weighted cost of capital")
header(ws, 3, ["Component", "Value", "Source / formula"])
wacc_rows = [
    ("Risk-free rate", rf / 100, "CNBC U.S. 10-year Treasury, September 21, 2026"),
    ("Equity risk premium", erp / 100, "Standard model assumption"),
    ("Levered beta", beta, "StockAnalysis five-year beta"),
    ("Cost of equity", cost_equity / 100, "Risk-free rate + beta × ERP"),
    ("Pre-tax cost of debt", pretax_debt_cost / 100, "Normalized estimate reflecting higher-rate acquisition financing"),
    ("Normalized tax rate", tax_rate / 100, "Statutory assumption; TTM GAAP effective rate is not meaningful with a pretax loss"),
    ("Market cap ($B)", MARKET_CAP, "StockAnalysis"),
    ("Total debt ($B)", DEBT, "StockAnalysis standardized total debt"),
    ("Equity weight", equity_weight, "Market cap / (market cap + debt)"),
    ("Debt weight", debt_weight, "Debt / (market cap + debt)"),
    ("After-tax debt cost", pretax_debt_cost / 100 * (1 - tax_rate / 100), "Kd × (1 − tax rate)"),
    ("WACC", wacc / 100, "E/V × Ke + D/V × Kd × (1 − t)"),
]
for row, values in enumerate(wacc_rows, 4):
    for col, value in enumerate(values, 1):
        cell = put(ws, row, col, value, bold=(col == 1), fill=(gold if values[0] == "WACC" else None))
        if col == 2 and values[0] not in {"Levered beta", "Market cap ($B)", "Total debt ($B)"}:
            cell.number_format = "0.00%"
widths(ws, [31, 20, 80, 14])

# 3. Scenarios
ws = wb.create_sheet("Scenarios")
title(ws, "Talen Energy — Forward P/E Scenarios", 6, "Five-year adjusted EPS framework; reported FCF multiples shown only as a leverage warning")
header(ws, 3, ["Metric", "Bear", "Base", "Bull", "Notes"])
scenario_rows = [
    ("Revenue CAGR (5Y)", *(SCENARIOS[k]["rev_cagr"] for k in ("Bear", "Base", "Bull")), "Applied from FY2027 visible $5.24B headline consensus"),
    ("Terminal revenue ($B)", *(SCENARIOS[k]["terminal_revenue"] for k in ("Bear", "Base", "Bull")), "Acquisitions, PJM prices, AWS ramp and additional contracting determine the path"),
    ("Terminal adjusted EPS", *(SCENARIOS[k]["terminal_eps"] for k in ("Bear", "Base", "Bull")), "Per-share outcome includes operating cash flow and repurchase effects"),
    ("Exit P/E", *(SCENARIOS[k]["multiple"] for k in ("Bear", "Base", "Bull")), "10x merchant/regulatory stress; 13x normalized; 14x infrastructure-like contracted mix"),
    ("Target price", *(SCENARIOS[k]["target"] for k in ("Bear", "Base", "Bull")), "Terminal adjusted EPS × exit P/E"),
    ("Upside / (downside)", *((SCENARIOS[k]["target"] / PRICE - 1) for k in ("Bear", "Base", "Bull")), "Versus $299.97 close"),
    ("Probability", *(SCENARIOS[k]["weight"] for k in ("Bear", "Base", "Bull")), "20% / 50% / 30%"),
    ("Weighted value/share", *(SCENARIOS[k]["target"] * SCENARIOS[k]["weight"] for k in ("Bear", "Base", "Bull")), "Contribution to probability-weighted fair value"),
]
for row, values in enumerate(scenario_rows, 4):
    for col, value in enumerate(values, 1):
        cell = put(ws, row, col, value, bold=(col == 1))
        if col in (2, 3, 4) and values[0] in {"Revenue CAGR (5Y)", "Upside / (downside)", "Probability"}:
            cell.number_format = "0.0%"
        elif col in (2, 3, 4) and values[0] in {"Terminal revenue ($B)", "Terminal adjusted EPS", "Target price", "Weighted value/share"}:
            cell.number_format = "$0.00"
put(ws, 13, 1, "Probability-weighted fair value", bold=True, fill=gold)
put(ws, 13, 2, weighted_value, bold=True, fill=gold).number_format = "$0.00"
put(ws, 14, 1, "Upside from current price", bold=True, fill=gold)
put(ws, 14, 2, weighted_value / PRICE - 1, bold=True, fill=gold).number_format = "0.0%"
put(ws, 16, 1, "Framework note", bold=True, fill=gray)
put(ws, 16, 2, "Net debt of $9.34B is 21.0× reported TTM FCF of $444M. The June acquisition close also makes trailing results structurally incomparable with management's $1.20B-$1.35B FY2026 adjusted FCF guidance. Forward adjusted P/E is therefore primary; leverage, basis and realized FCF are the cross-checks.", fill=gray)
ws.merge_cells("B16:E16")
widths(ws, [31, 18, 18, 18, 86, 14])

# 4. Actuals Source Audit
ws = wb.create_sheet("Actuals Source Audit")
title(ws, "Talen Energy — Actuals Source Audit", 5, "All financial statement figures are USD millions unless noted")
header(ws, 3, ["Data point", "Value", "Source URL", "Source date", "Notes"])
sa = "https://stockanalysis.com/stocks/tln"
audit = [
    ("Stock price", "$299.97", f"{sa}/", "2026-09-21", "Official close; after-hours excluded"),
    ("Market cap", "$14.37B", f"{sa}/statistics/", "2026-09-21", "Price-sensitive daily statistic"),
    ("Enterprise value", "$23.71B", f"{sa}/statistics/", "2026-09-21", "EV less MC implies $9.34B net debt"),
    ("Shares outstanding", "47.91M", f"{sa}/statistics/", "2026-09-21", "Down 7.07% YoY; acquisition issuance and buybacks both affect the path"),
    ("Revenue TTM", "$3,741M", f"{sa}/financials/", "2026-06-30", "+75.8%; includes acquisition timing effects"),
    ("Operating income TTM", "$153M", f"{sa}/statistics/", "2026-06-30", "4.09% margin"),
    ("Net income TTM", "-$185M", f"{sa}/financials/", "2026-06-30", "GAAP; below-the-line and transaction items make adjusted cash flow more relevant"),
    ("Cash and investments", "$232M", f"{sa}/financials/balance-sheet/", "2026-06-30", "$231M cash plus $1M securities"),
    ("Total debt", "$9,574M", f"{sa}/financials/balance-sheet/", "2026-06-30", "Rose from $6,833M at FY2025 and $3,004M at FY2024"),
    ("Common equity", "$1,616M", f"{sa}/financials/balance-sheet/", "2026-06-30", "Retained deficit of $905M"),
    ("Operating cash flow TTM", "$796M", f"{sa}/financials/cash-flow-statement/", "2026-06-30", "Includes $568M stock-based compensation and working-capital effects"),
    ("Capital expenditures TTM", "-$352M", f"{sa}/financials/cash-flow-statement/", "2026-06-30", "Provider FCF uses total capex including nuclear fuel expenditures"),
    ("Reported free cash flow TTM", "$444M", f"{sa}/financials/cash-flow-statement/", "2026-06-30", "11.87% margin; differs from management adjusted FCF"),
    ("Cash acquisitions TTM", "-$6,361M", f"{sa}/financials/cash-flow-statement/", "2026-06-30", "Primary reason investing cash flow is far below capex"),
    ("Repurchases TTM", "-$438M", f"{sa}/financials/cash-flow-statement/", "2026-06-30", "Supports per-share cash flow but competes with deleveraging"),
    ("FY2026 revenue consensus", "$4.53B", f"{sa}/forecast/", "2026-09-18", "10 analysts; +72.5%"),
    ("FY2026 adjusted EPS consensus", "$22.26", f"{sa}/forecast/", "2026-09-18", "Non-GAAP adjusted diluted; range $17.43-$25.96"),
    ("FY2027 revenue headline", "$5.24B", f"{sa}/forecast/", "2026-09-18", "Public headline; detailed table gated"),
    ("FY2027 adjusted EPS headline", "$30.75", f"{sa}/forecast/", "2026-09-18", "Public headline; detailed table gated"),
    ("Average analyst target", "$459.94", f"{sa}/forecast/", "2026-09-18", "17 analysts; range $307-$560"),
    ("Beta", "1.63", f"{sa}/statistics/", "2026-09-21", "Five-year beta"),
    ("10Y Treasury", "4.957%", "https://www.cnbc.com/2026/09/21/treasury-yields-government-bonds.html", "2026-09-21", "Benchmark yield"),
    ("Q2 adjusted EBITDA", "$374M", f"{sa}/transcripts/660922-q2-2026/", "2026-08-05", "Management non-GAAP figure"),
    ("Q2 adjusted FCF", "$212M", f"{sa}/transcripts/660922-q2-2026/", "2026-08-05", "Management non-GAAP figure"),
    ("FY2026 adjusted EBITDA guidance", "$2,025M-$2,225M", f"{sa}/transcripts/660922-q2-2026/", "2026-08-05", "Raised after Cornerstone close"),
    ("FY2026 adjusted FCF guidance", "$1,200M-$1,350M", f"{sa}/transcripts/660922-q2-2026/", "2026-08-05", "Management non-GAAP; includes financing impacts"),
    ("Next earnings date", "Nov. 4, 2026", f"{sa}/statistics/", "2026-09-21", "Estimated, after market close"),
]
for row, values in enumerate(audit, 4):
    for col, value in enumerate(values, 1):
        put(ws, row, col, value, bold=(col == 1))
widths(ws, [31, 23, 88, 16, 72])

# 5. Questions
ws = wb.create_sheet("Questions")
title(ws, "Talen Energy — Open Questions", 4, "Items that can change the valuation or risk assessment")
header(ws, 3, ["#", "Question", "Why it matters", "Best evidence / next check"])
questions = [
    (1, "How quickly can net leverage return to management's 3.5× target after the Cornerstone acquisition?", "Debt rose to $9.57B; cash returns and M&A compete with deleveraging.", "Q3 leverage bridge, debt repayment and covenant disclosures."),
    (2, "What reconciles $444M reported TTM FCF to $1.20B-$1.35B FY2026 adjusted FCF guidance?", "The equity thesis depends on adjusted cash becoming cash available for repurchases.", "Management reconciliation by hedges, working capital, nuclear fuel and transaction items."),
    (3, "How much of the revenue and FCF step-up is acquisition-driven versus organic?", "The June close makes trailing growth rates non-comparable.", "Pro forma same-asset revenue, EBITDA and cash flow."),
    (4, "What is the sustainable PPL-to-West Hub basis after transmission work ends?", "Management says each $1 basis improvement adds about $1 of adjusted FCF per share.", "Forward curves, realized basis and transmission completion schedule."),
    (5, "How much of the merchant generation portfolio can be contracted without giving away upside?", "Long-term PPAs lower risk and cost of capital but cap merchant optionality.", "Contracted gross-margin mix and PPA terms."),
    (6, "Will the AWS ramp reach full build-out in 2028-2030 as planned?", "The existing ~2 GW contract is expected to lift contracted gross margin from 10% toward 35%.", "Campus load milestones and contract economics."),
    (7, "Can an additional ~2 GW of long-term contracts be signed at attractive returns?", "Management's illustrative 60% contracted gross-margin mix depends on this execution.", "Signed PPAs, counterparty credit and required new capacity."),
    (8, "How will PJM's RBP and IRAS rules affect existing assets, new capacity and data-center load?", "Regulatory design can change capacity revenue, curtailment and project economics.", "Final FERC orders and company participation."),
    (9, "How durable are PJM capacity prices after temporary auction caps roll off?", "Capacity clearing at caps supports current cash flow but may attract policy intervention and supply.", "Auction outcomes and uncapped clearing estimates."),
    (10, "What drove $6.36B of TTM cash acquisitions, and what synergies are embedded in guidance?", "Purchase accounting and integration quality determine whether debt-funded growth creates equity value.", "Acquisition purchase-price allocations and pro forma synergy bridge."),
    (11, "Why was TTM stock-based compensation $568M versus only $33M in FY2024?", "SBC exceeded reported FCF and can overstate cash conversion if it is transaction-related but recurring dilution is ignored.", "10-Q SBC footnote and award vesting schedule."),
    (12, "Can buybacks remain accretive after acquisition-related equity issuance?", "Management targets returning 70% of adjusted FCF, but timing and repurchase price drive per-share value.", "Quarterly share-count reconciliation and average repurchase price."),
    (13, "What portion of capex is maintenance, nuclear fuel, growth uprates, batteries and peakers?", "Recurring owner earnings cannot be estimated without separating maintenance and growth investment.", "Annual capex and nuclear fuel guidance."),
    (14, "What are Susquehanna outage, decommissioning and nuclear regulatory obligations?", "Nuclear reliability is a major cash-flow driver and low-frequency risk.", "Outage schedule, NRC disclosures and decommissioning trust funding."),
    (15, "What must Q3 prove on November 4?", "The quarter should validate Cornerstone integration, FY2027 guidance and the adjusted FCF/share bridge.", "Q3 release and call."),
]
for row, values in enumerate(questions, 4):
    for col, value in enumerate(values, 1):
        put(ws, row, col, value, bold=(col == 1))
widths(ws, [7, 80, 68, 60])

# 6. Sources
ws = wb.create_sheet("Sources")
title(ws, "Talen Energy — Sources", 4, "Public sources accessed September 22, 2026")
header(ws, 3, ["#", "Source", "URL", "Use"])
sources = [
    (1, "StockAnalysis overview", f"{sa}/", "Price, identity, market cap and headline results"),
    (2, "StockAnalysis financials", f"{sa}/financials/", "Income statement, segments, margins and historical overview"),
    (3, "StockAnalysis balance sheet", f"{sa}/financials/balance-sheet/", "Cash, debt, assets, liabilities and equity"),
    (4, "StockAnalysis cash flow", f"{sa}/financials/cash-flow-statement/", "OCF, capex, FCF, acquisitions, financing and buybacks"),
    (5, "StockAnalysis statistics", f"{sa}/statistics/", "Valuation, beta, EV, leverage and analyst summary"),
    (6, "StockAnalysis forecast", f"{sa}/forecast/", "Adjusted EPS, revenue consensus and price targets"),
    (7, "StockAnalysis profile", f"{sa}/company/", "Company identity, assets, management and filings"),
    (8, "Talen Q2 2026 transcript", f"{sa}/transcripts/660922-q2-2026/", "Strategy, guidance, acquisitions, PJM, PPAs and capital returns"),
    (9, "Talen Q2 2026 release", "https://stockanalysis.com/filings/TLN/3495601/", "Quarterly adjusted results and guidance"),
    (10, "CNBC Treasury yields", "https://www.cnbc.com/2026/09/21/treasury-yields-government-bonds.html", "Risk-free rate"),
    (11, "Talen investor relations", "https://ir.talenenergy.com/", "Primary materials and future updates"),
]
for row, values in enumerate(sources, 4):
    for col, value in enumerate(values, 1):
        put(ws, row, col, value, bold=(col == 1))
widths(ws, [7, 36, 100, 68])

for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions

wb.save(OUT)
check = load_workbook(OUT, data_only=False)
expected = ["Valuation", "WACC", "Scenarios", "Actuals Source Audit", "Questions", "Sources"]
assert check.sheetnames == expected, check.sheetnames
assert check["Scenarios"]["B13"].value == weighted_value
print(f"WACC: {wacc:.2f}%")
for name, case in SCENARIOS.items():
    print(f"{name}: ${case['target']:.2f} ({case['target'] / PRICE - 1:.1%})")
print(f"Probability-weighted fair value: ${weighted_value:.2f} ({weighted_value / PRICE - 1:.1%})")
print(f"Wrote and verified {OUT}")
