#!/usr/bin/env python3
"""Build the Chord Energy six-sheet valuation workbook."""
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUT = Path(__file__).with_name("2026-09-16 Chord Energy Model.xlsx")
PRICE = 157.00
SHARES = 54.70  # millions
MARKET_CAP = 8.59  # billions
DEBT = 1.50  # billions
CASH = 0.612  # billions
NET_DEBT = DEBT - CASH
ENTERPRISE_VALUE = MARKET_CAP + NET_DEBT
TTM_REVENUE = 5.96  # billions
TTM_EBITDA = 2.74  # billions
TTM_FCF = 1.19  # billions
TTM_EPS = 14.81
FY26_REVENUE = 6.68  # S&P Global consensus, billions
FY26_EPS = 20.07  # adjusted diluted consensus

SCENARIOS = {
    "Bear": {"rev_cagr": -0.02, "fcf_margin": 0.12, "multiple": 6.0, "weight": 0.20},
    "Base": {"rev_cagr": 0.00, "fcf_margin": 0.18, "multiple": 8.0, "weight": 0.55},
    "Bull": {"rev_cagr": 0.025, "fcf_margin": 0.22, "multiple": 10.0, "weight": 0.25},
}
for case in SCENARIOS.values():
    case["terminal_revenue"] = FY26_REVENUE * (1 + case["rev_cagr"]) ** 5
    case["terminal_fcf"] = case["terminal_revenue"] * case["fcf_margin"]
    case["implied_ev"] = case["terminal_fcf"] * case["multiple"]
    case["target"] = (case["implied_ev"] - NET_DEBT) * 1000 / SHARES
weighted_value = sum(case["target"] * case["weight"] for case in SCENARIOS.values())

assert SCENARIOS["Bear"]["target"] < PRICE
assert abs(SCENARIOS["Base"]["target"] - 168.27) / 168.27 < 0.20
assert 50 < weighted_value < 350

rf = 5.04
erp = 5.00
beta = 0.39
cost_equity = rf + beta * erp
pretax_debt_cost = 7.00
tax_rate = 23.4
equity_weight = MARKET_CAP / (MARKET_CAP + DEBT)
debt_weight = 1 - equity_weight
wacc = equity_weight * cost_equity + debt_weight * pretax_debt_cost * (1 - tax_rate / 100)

wb = Workbook()
navy, blue, gold, gray, white = "17365D", "D9EAF7", "D8B34B", "E7E6E6", "FFFFFF"
thin = Side(style="thin", color="A6A6A6")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def put(ws, row, col, value, *, bold=False, fill=None, color="000000", size=10):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(name="Aptos", size=size, bold=bold, color=color)
    cell.fill = PatternFill("solid", fgColor=fill) if fill else PatternFill(fill_type=None)
    cell.border = border
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    return cell


def title(ws, text, end_col=5, subtitle=None):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    put(ws, 1, 1, text, bold=True, fill=navy, color=white, size=15)
    ws.row_dimensions[1].height = 28
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
        put(ws, 2, 1, subtitle, bold=True, fill=blue)


def header(ws, row, values):
    for col, value in enumerate(values, 1):
        put(ws, row, col, value, bold=True, fill=navy, color=white)


def widths(ws, values):
    for col, width in enumerate(values, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A3"


# 1. Valuation
ws = wb.active
ws.title = "Valuation"
title(ws, "Chord Energy (CHRD) — Valuation", 4, "Quote date: September 15, 2026 | Model date: September 16, 2026")
facts = [
    ("Company", "Chord Energy Corporation", "Independent E&P concentrated in the oil-rich Williston Basin"),
    ("Ticker", "NASDAQ: CHRD", "Energy / Oil & Gas Exploration & Production"),
    ("Price", PRICE, "September 15 close"),
    ("Shares outstanding (M)", SHARES, "Current StockAnalysis overview; Q2 filing reported 55.2M at June 30"),
    ("Market capitalization ($B)", MARKET_CAP, "Price × current shares"),
    ("Enterprise value ($B)", ENTERPRISE_VALUE, "Market cap plus debt less cash"),
    ("Net debt ($B)", NET_DEBT, "June 30 debt less cash"),
    ("Primary valuation lens", "Normalized FCF multiple", "Commodity scenarios require explicit oil-price and margin sensitivity"),
    ("Stance", "Watch / Positive", "Strong balance sheet and returns; current price already discounts much of FY2026 strength"),
]
header(ws, 3, ["Field", "Value", "Comment"])
for row, values in enumerate(facts, 4):
    for col, value in enumerate(values, 1):
        put(ws, row, col, value, bold=(col == 1))

header(ws, 15, ["Valuation metric", "Value", "Interpretation"])
metrics = [
    ("Trailing P/E", PRICE / TTM_EPS, "GAAP TTM; derivative marks and impairments create quarter-to-quarter noise"),
    ("Forward P/E", PRICE / FY26_EPS, "FY2026 adjusted diluted consensus"),
    ("P/S", MARKET_CAP / TTM_REVENUE, "Equity value relative to commodity-sensitive revenue"),
    ("P/FCF", MARKET_CAP / TTM_FCF, "Current FCF is helped by a favorable commodity environment"),
    ("EV/FCF", ENTERPRISE_VALUE / TTM_FCF, "Primary current cash-flow cross-check"),
    ("EV/Sales", ENTERPRISE_VALUE / TTM_REVENUE, "Low multiple reflects commodity cyclicality"),
    ("EV/EBITDA", ENTERPRISE_VALUE / TTM_EBITDA, "Secondary E&P cross-check"),
    ("Net debt / FCF", NET_DEBT / TTM_FCF, "Balance sheet has substantial commodity-cycle resilience"),
    ("FCF yield", TTM_FCF / MARKET_CAP, "$1.19B TTM FCF / $8.59B market cap"),
]
for row, values in enumerate(metrics, 16):
    for col, value in enumerate(values, 1):
        cell = put(ws, row, col, value, bold=(col == 1))
        if values[0] == "FCF yield" and col == 2:
            cell.number_format = "0.0%"
widths(ws, [28, 24, 78, 14])

# 2. WACC
ws = wb.create_sheet("WACC")
title(ws, "Chord Energy — WACC", 4, "CAPM and debt-weighted cost of capital; scenario multiples remain the primary valuation tool")
header(ws, 3, ["Component", "Value", "Source / formula"])
wacc_rows = [
    ("Risk-free rate", rf / 100, "CNBC U.S. 10-year Treasury intraday high, September 16, 2026"),
    ("Equity risk premium", erp / 100, "Standard model assumption"),
    ("Levered beta", beta, "StockAnalysis five-year beta"),
    ("Cost of equity", cost_equity / 100, "Risk-free + beta × ERP"),
    ("Pre-tax cost of debt", pretax_debt_cost / 100, "Normalized estimate; Q2 cash interest annualized / debt is approximately 6.9%"),
    ("Tax rate", tax_rate / 100, "Q2 adjustment-item tax rate"),
    ("Market cap ($B)", MARKET_CAP, "September 15 quote"),
    ("Total debt ($B)", DEBT, "June 30 balance sheet"),
    ("Equity weight", equity_weight, "Market cap / (market cap + debt)"),
    ("Debt weight", debt_weight, "Debt / (market cap + debt)"),
    ("After-tax debt cost", pretax_debt_cost / 100 * (1 - tax_rate / 100), "Kd × (1 − tax rate)"),
    ("WACC", wacc / 100, "E/V × Ke + D/V × Kd × (1 − t)"),
]
for row, values in enumerate(wacc_rows, 4):
    for col, value in enumerate(values, 1):
        cell = put(ws, row, col, value, bold=(col == 1 or values[0] == "WACC"), fill=(gold if values[0] == "WACC" else None))
        if col == 2 and values[0] not in {"Levered beta", "Market cap ($B)", "Total debt ($B)"}:
            cell.number_format = "0.00%"
widths(ws, [30, 20, 76, 14])

# 3. Scenarios
ws = wb.create_sheet("Scenarios")
title(ws, "Chord Energy — Normalized FCF Scenarios", 6, "Five-year commodity-cycle framework; all financial figures in $B except per-share values")
header(ws, 3, ["Metric", "Bear", "Base", "Bull", "Notes"])
scenario_rows = [
    ("Revenue CAGR (5Y)", *(SCENARIOS[k]["rev_cagr"] for k in ("Bear", "Base", "Bull")), "From FY2026 consensus; price realization dominates reported revenue"),
    ("Terminal revenue ($B)", *(SCENARIOS[k]["terminal_revenue"] for k in ("Bear", "Base", "Bull")), "FY2026 consensus compounded five years"),
    ("Adjusted FCF margin", *(SCENARIOS[k]["fcf_margin"] for k in ("Bear", "Base", "Bull")), "Commodity-price and capital-efficiency sensitivity"),
    ("Terminal FCF ($B)", *(SCENARIOS[k]["terminal_fcf"] for k in ("Bear", "Base", "Bull")), "Terminal revenue × adjusted FCF margin"),
    ("Exit FCF multiple", *(SCENARIOS[k]["multiple"] for k in ("Bear", "Base", "Bull")), "6x stress / 8x normalized / 10x durable execution"),
    ("Implied EV ($B)", *(SCENARIOS[k]["implied_ev"] for k in ("Bear", "Base", "Bull")), "Terminal FCF × exit multiple"),
    ("Less net debt ($B)", *(NET_DEBT for _ in range(3)), "June 30 debt less cash; no assumed future deleveraging"),
    ("Shares (M)", *(SHARES for _ in range(3)), "Current shares; buybacks excluded for conservatism"),
    ("Target price", *(SCENARIOS[k]["target"] for k in ("Bear", "Base", "Bull")), "(Implied EV − net debt) / shares"),
    ("Upside / (downside)", *((SCENARIOS[k]["target"] / PRICE - 1) for k in ("Bear", "Base", "Bull")), "Versus $157.00"),
    ("Probability", *(SCENARIOS[k]["weight"] for k in ("Bear", "Base", "Bull")), "20% / 55% / 25%"),
    ("Weighted value/share", *(SCENARIOS[k]["target"] * SCENARIOS[k]["weight"] for k in ("Bear", "Base", "Bull")), "Contribution to fair value"),
]
for row, values in enumerate(scenario_rows, 4):
    for col, value in enumerate(values, 1):
        cell = put(ws, row, col, value, bold=(col == 1))
        if col in (2, 3, 4) and values[0] in {"Revenue CAGR (5Y)", "Adjusted FCF margin", "Upside / (downside)", "Probability"}:
            cell.number_format = "0.0%"
        elif col in (2, 3, 4) and values[0] in {"Terminal revenue ($B)", "Terminal FCF ($B)", "Implied EV ($B)", "Less net debt ($B)", "Target price", "Weighted value/share"}:
            cell.number_format = "$0.00"
put(ws, 17, 1, "Probability-weighted fair value", bold=True, fill=gold)
put(ws, 17, 2, weighted_value, bold=True, fill=gold).number_format = "$0.00"
put(ws, 18, 1, "Upside from current price", bold=True, fill=gold)
put(ws, 18, 2, weighted_value / PRICE - 1, bold=True, fill=gold).number_format = "0.0%"
widths(ws, [30, 18, 18, 18, 82, 14])

# 4. Actuals Source Audit
ws = wb.create_sheet("Actuals Source Audit")
title(ws, "Chord Energy — Actuals Source Audit", 5, "Dollar figures are USD; dated source snapshots are not silently blended")
header(ws, 3, ["Data point", "Value", "Source URL", "Source date", "Notes"])
sa = "https://stockanalysis.com/stocks/chrd"
release = "https://www.prnewswire.com/news-releases/chord-energy-reports-second-quarter-2026-financial-and-operating-results-declares-base-dividend-and-updates-2026-outlook-302844081.html"
audit = [
    ("Stock price", "$157.00", f"{sa}/", "2026-09-15", "Official close"),
    ("Market cap", "$8.59B", f"{sa}/", "2026-09-15", "Current overview"),
    ("Enterprise value", "$9.48B", "Calculated from market cap plus debt less cash", "2026-09-16", "Uses June 30 balance sheet"),
    ("Shares outstanding", "54.70M", f"{sa}/", "2026-09-15", "Q2 filing reported 55.20M at June 30"),
    ("Revenue TTM", "$5.96B", f"{sa}/", "2026-09-15", "Current overview"),
    ("Net income TTM", "$840.73M", f"{sa}/", "2026-09-15", "GAAP"),
    ("EPS TTM", "$14.81", f"{sa}/", "2026-09-15", "GAAP provider figure"),
    ("EBITDA TTM", "$2.74B", f"{sa}/statistics/", "2026-06-30", "Latest visible provider financial snapshot"),
    ("Operating cash flow TTM", "$2.59B", f"{sa}/statistics/", "2026-06-30", "Provider standardized figure"),
    ("Capital expenditures TTM", "$1.40B", f"{sa}/statistics/", "2026-06-30", "Provider standardized figure"),
    ("Free cash flow TTM", "$1.19B", f"{sa}/statistics/", "2026-06-30", "OCF less capex"),
    ("Cash", "$611.57M", release, "2026-06-30", "Company-reported cash and equivalents"),
    ("Total debt", "$1.50B", release, "2026-06-30", "No revolver borrowing; senior notes"),
    ("Stockholders' equity", "$8.36B", release, "2026-06-30", "Company-reported"),
    ("Q2 total revenue", "$2.17B", release, "2026-06-30", "Includes purchased oil and gas sales"),
    ("Q2 net income", "$525.2M", release, "2026-06-30", "GAAP"),
    ("Q2 adjusted FCF", "$413.4M", release, "2026-06-30", "Non-GAAP; excludes reimbursable non-op capex"),
    ("FY2026 adjusted EBITDA guidance", "$3.0B", release, "2026-08-05", "Assumes $75 WTI and $3 Henry Hub in 2H26"),
    ("FY2026 adjusted FCF guidance", "$1.3B", release, "2026-08-05", "Same commodity assumptions"),
    ("FY2026 revenue consensus", "$6.68B", f"{sa}/forecast/", "2026-09-16", "S&P Global; six visible analysts"),
    ("FY2026 adjusted EPS consensus", "$20.07", f"{sa}/forecast/", "2026-09-16", "Adjusted diluted provider estimate"),
    ("Average analyst target", "$168.27", f"{sa}/forecast/", "2026-09-16", "16 analysts; range $130-$217"),
    ("Beta", "0.39", f"{sa}/", "2026-09-15", "Five-year beta"),
    ("10Y Treasury", "5.04%", "https://www.cnbc.com/quotes/US10Y", "2026-09-16", "Intraday high used conservatively"),
    ("Next earnings", "November 3, 2026", f"{sa}/", "2026-09-15", "Expected date"),
]
for row, values in enumerate(audit, 4):
    for col, value in enumerate(values, 1):
        put(ws, row, col, value, bold=(col == 1))
widths(ws, [31, 21, 88, 16, 68])

# 5. Questions
ws = wb.create_sheet("Questions")
title(ws, "Chord Energy — Open Questions", 4, "Items that can materially change normalized FCF or the terminal multiple")
header(ws, 3, ["#", "Question", "Why it matters", "Best evidence / next check"])
questions = [
    (1, "How much of the FY2026 production and cost improvement is durable rather than a function of accelerated TIL timing?", "Q2 benefited from completions pulled forward, while Q4 oil volume is guided lower.", "Q3/Q4 production and capex bridge."),
    (2, "Do 4-mile laterals sustain type-curve economics after a full production history?", "Longer laterals can lower unit costs, but early performance is not a mature decline curve.", "12- and 18-month well productivity by vintage."),
    (3, "What WTI price supports the base dividend and 75% return-of-capital target after maintenance capex?", "Shareholder returns are the central equity proposition.", "Price-sensitivity table and maintenance capital disclosure."),
    (4, "How much of $1.3B FY2026 adjusted FCF guidance depends on hedge gains or losses?", "Derivative timing can obscure underlying field economics.", "Realized/unrealized hedge reconciliation."),
    (5, "Why did 2Q26 production taxes slightly exceed guidance?", "Recurring per-barrel leakage reduces cash margins.", "Tax rate by commodity and jurisdiction."),
    (6, "Will the chemical workover program improve decline rates enough to offset higher LOE?", "Management raised LOE guidance partly to fund production enhancement.", "Incremental barrels, cost per barrel and payout period."),
    (7, "How much inventory remains at current spacing and return thresholds?", "Williston concentration makes inventory depth central to terminal value.", "High-return location count at $60/$70/$80 WTI."),
    (8, "What is the acquisition hurdle rate after the Enerplus combination and recent bolt-ons?", "Acquisitions drove scale but can dilute per-share returns at cycle peaks.", "Deal-level synergy and return scorecard."),
    (9, "Why did total assets and debt jump sharply after FY2023?", "The Enerplus combination changed comparability, share count and capital structure.", "Purchase accounting and pro-forma financials."),
    (10, "Can the company hold cash G&A below guidance as integration benefits mature?", "Q2 cash G&A beat guidance and supports the synergy thesis.", "Quarterly cash G&A and merger-cost run-off."),
    (11, "What portion of purchased oil and gas sales is low-margin and should be excluded from organic revenue comparisons?", "Gross presentation inflates revenue without equivalent economic value.", "Purchased-sales gross margin and volumes."),
    (12, "What is normalized effective tax rate after derivative and impairment volatility?", "Tax expense can diverge sharply from pretax income in noisy quarters.", "Cash-tax guidance and deferred-tax roll-forward."),
    (13, "Can buybacks continue near $150-$160 without reducing per-share returns?", "Q2 repurchases at $133.47 were clearly more accretive than purchases near the current high.", "Repurchase price, authorization and NAV sensitivity."),
    (14, "What customer, pipeline and rail concentration could impair realized pricing?", "Single-basin concentration creates takeaway and counterparty risk.", "Top purchasers and firm transportation commitments."),
    (15, "How sensitive are reserves and asset retirement obligations to lower long-term commodity prices?", "Reserve revisions affect both NAV and future DD&A.", "Year-end reserve report and standardized measure."),
    (16, "What will the November 3 earnings release prove?", "The key test is whether production optimization offsets fewer Q4 TILs while preserving FCF conversion.", "Q3 actuals, Q4 guide and preliminary 2027 capital plan."),
]
for row, values in enumerate(questions, 4):
    for col, value in enumerate(values, 1):
        put(ws, row, col, value, bold=(col == 1))
widths(ws, [7, 80, 68, 60])

# 6. Sources
ws = wb.create_sheet("Sources")
title(ws, "Chord Energy — Sources", 4, "Public sources accessed September 16, 2026")
header(ws, 3, ["#", "Source", "URL", "Use"])
sources = [
    (1, "StockAnalysis overview", f"{sa}/", "Quote, company identity and current headline financials"),
    (2, "StockAnalysis income statement", f"{sa}/financials/", "Historical income, margins, shares and EBITDA"),
    (3, "StockAnalysis balance sheet", f"{sa}/financials/balance-sheet/", "Historical cash, debt, assets and equity"),
    (4, "StockAnalysis cash flow", f"{sa}/financials/cash-flow-statement/", "Historical OCF, capex, acquisitions, dividends and repurchases"),
    (5, "StockAnalysis statistics", f"{sa}/statistics/", "Valuation, leverage and standardized TTM financials"),
    (6, "StockAnalysis forecast", f"{sa}/forecast/", "S&P Global revenue/EPS forecasts and analyst targets"),
    (7, "StockAnalysis company profile", f"{sa}/company/", "Business description, leadership and listing details"),
    (8, "Chord Q2 2026 release", release, "Primary financials, operations, guidance, liquidity and capital returns"),
    (9, "Chord investor relations", "https://ir.chordenergy.com/", "Future primary-source updates and presentations"),
    (10, "CNBC U.S. 10-year Treasury", "https://www.cnbc.com/quotes/US10Y", "Risk-free rate"),
]
for row, values in enumerate(sources, 4):
    for col, value in enumerate(values, 1):
        put(ws, row, col, value, bold=(col == 1))
widths(ws, [7, 34, 100, 72])

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and cell.row > 2:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

wb.save(OUT)
check = load_workbook(OUT, data_only=False)
assert check.sheetnames == ["Valuation", "WACC", "Scenarios", "Actuals Source Audit", "Questions", "Sources"]
assert abs(check["Scenarios"]["B17"].value - weighted_value) < 1e-9
print(f"WACC: {wacc:.2f}%")
print("Targets:", {k: round(v["target"], 2) for k, v in SCENARIOS.items()})
print(f"Probability-weighted FV: ${weighted_value:.2f} ({weighted_value / PRICE - 1:+.1%})")
print(f"Saved {OUT}")
