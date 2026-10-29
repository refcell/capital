#!/usr/bin/env python3
"""Build the 2026-10-29 Xencor valuation workbook (cash NAV + risk-adjusted pipeline)."""

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


OUT = Path(__file__).with_name("2026-10-29 Xencor Model.xlsx")
PRICE = 29.99  # Oct. 28, 2026 close
SHARES = 74.35  # millions, StockAnalysis shares outstanding (filing date)
MARKET_CAP = 2230.0  # $ millions, StockAnalysis
PROVIDER_EV = 1830.0  # $ millions, StockAnalysis
CASH_INV = 564.99  # cash + marketable debt and equity securities, Jun. 30, 2026 ($M)
CASH_DEBT_SEC = 486.4  # company-defined cash, equivalents and marketable debt securities
EQUITY_SEC = 78.6  # marketable equity securities (partner stock), Jun. 30, 2026
ROYALTY_LIAB = 104.9  # non-recourse liability for sale of future royalties (OMERS)
LEASES = 65.2  # operating lease liabilities
TOTAL_DEBT = 170.11  # provider total debt = royalty liability + leases
NET_CASH = CASH_INV - TOTAL_DEBT  # $394.9M, matches provider net cash
ALEXION_2ND = 52.5  # second settlement installment due ~Aug. 2027
REV_TTM = 104.97
REV_FY25 = 125.58
REV_FY26_CONS = 148.84
REV_FY27_CONS = 121.03
EPS_FY26_CONS = -3.07
EPS_FY27_CONS = -3.54
OP_INC_TTM = -216.08
NI_TTM = -163.33
OCF_TTM = -205.81
FCF_TTM = -208.29
SBC_TTM = 39.03
BVPS = 6.77
ANALYST_PT = 40.25
YE26_CASH_GUIDE = 430.0  # midpoint of $420-$440M guidance (cash + marketable debt securities)

NAVY = "17365D"
BLUE = "D9EAF7"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
WHITE = "FFFFFF"
thin = Side(style="thin", color="B7B7B7")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def put(ws, row, col, value, *, bold=False, fill=None, fmt=None, wrap=False):
    cell = ws.cell(row=row, column=col, value=value)
    cell.border = border
    cell.font = Font(bold=bold, color=WHITE if fill == NAVY else "000000")
    cell.alignment = Alignment(vertical="top", wrap_text=wrap)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if fmt:
        cell.number_format = fmt
    return cell


def title(ws, text, end_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    cell = ws.cell(1, 1, text)
    cell.font = Font(size=15, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 24


def headers(ws, row, values):
    for col, value in enumerate(values, 1):
        put(ws, row, col, value, bold=True, fill=BLUE, wrap=True)


# WACC
rf = 0.0502
erp = 0.05
beta = 0.90
ke = rf + beta * erp
cost_debt = 0.15  # blended: OMERS royalty liability accretes ~21% annualized (H1 2026), leases ~6%
tax = 0.0  # NOL-sheltered; no tax shield
equity_weight = MARKET_CAP / (MARKET_CAP + TOTAL_DEBT)
debt_weight = 1 - equity_weight
wacc = equity_weight * ke + debt_weight * cost_debt * (1 - tax)
rnpv_rate = 0.12  # selected discount rate embedded in pipeline values (clinical-stage biotech)

# Scenario inputs, $ millions. XmAb819 rNPV = peak sales x probability of success x NPV/peak multiple.
scen = {
    "Bear": {"royalty": 150.0, "peak819": 800.0, "pos819": 0.20, "mult819": 2.0, "v942": 50.0, "other": 50.0, "corp": -250.0, "dil": 1.25, "w": 0.25},
    "Base": {"royalty": 200.0, "peak819": 2000.0, "pos819": 0.40, "mult819": 2.5, "v942": 300.0, "other": 150.0, "corp": -300.0, "dil": 1.15, "w": 0.50},
    "Bull": {"royalty": 280.0, "peak819": 3000.0, "pos819": 0.45, "mult819": 2.5, "v942": 900.0, "other": 350.0, "corp": -250.0, "dil": 1.05, "w": 0.25},
}
for s in scen.values():
    s["net_cash"] = NET_CASH
    s["alexion"] = ALEXION_2ND / (1 + rnpv_rate) ** 0.8
    s["v819"] = s["peak819"] * s["pos819"] * s["mult819"]
    s["equity"] = s["net_cash"] + s["alexion"] + s["royalty"] + s["v819"] + s["v942"] + s["other"] + s["corp"]
    s["shares"] = SHARES * s["dil"]
    s["target"] = s["equity"] / s["shares"]
    s["pipeline_ps"] = (s["v819"] + s["v942"] + s["other"]) / s["shares"]
weighted_fv = sum(s["target"] * s["w"] for s in scen.values())
net_cash_ps = NET_CASH / SHARES
premium_ps = PRICE - net_cash_ps

wb = Workbook()
ws = wb.active
ws.title = "Valuation"
title(ws, "Xencor, Inc. (NASDAQ: XNCR) — Valuation Summary", 4)
summary = [
    ("Company", "Xencor, Inc. (clinical-stage engineered-antibody biotech, Pasadena CA)", "Model date", "2026-10-29"),
    ("Quote", "$29.99 at Oct. 28, 2026 close (52-wk $10.38-$33.60)", "Shares outstanding", "74.35M (74.33M at Jun. 30, 2026 balance sheet)"),
    ("Market capitalization", "$2.23B", "Enterprise value", f"${PROVIDER_EV / 1000:.2f}B (MC − ${NET_CASH:.0f}M net cash)"),
    ("Net cash (Jun. 30, 2026)", f"${NET_CASH:.1f}M = ${CASH_INV:.0f}M cash & investments − ${ROYALTY_LIAB:.1f}M OMERS royalty liability − ${LEASES:.1f}M leases", "Liquidity guide", "YE2026 cash + marketable debt securities $420-$440M; runway through 2028"),
    ("Primary lens", "Cash NAV + risk-adjusted pipeline NPV (sum of parts); P/E and FCF multiples not meaningful", "Stance", "Watch"),
]
for r, row in enumerate(summary, 3):
    for c, value in enumerate(row, 1):
        put(ws, r, c, value, bold=c in (1, 3), fill=YELLOW if c in (1, 3) else None, wrap=True)

headers(ws, 10, ["Metric", "Current / Forward", "Interpretation", "Source date"])
metrics = [
    ("Trailing P/E", "n/m", f"TTM EPS −$2.18; net loss ${NI_TTM:.1f}M flattered by ~$64M investment gains", "2026-06-30"),
    ("Forward P/E (FY2026E / FY2027E)", "n/m", f"Consensus EPS ${EPS_FY26_CONS:.2f} / ${EPS_FY27_CONS:.2f} (non-GAAP provider data, 13 analysts)", "2026-10-27"),
    ("P/S (TTM)", f"{MARKET_CAP / REV_TTM:.1f}x", f"TTM revenue ${REV_TTM:.1f}M is royalties and partner milestones, not product sales", "2026-06-30"),
    ("EV / Sales (TTM)", f"{PROVIDER_EV / REV_TTM:.1f}x", "Revenue includes a $27.6M one-time Ultomiris catch-up royalty in Q2 2026", "2026-06-30"),
    ("EV / FY2026E revenue", f"{PROVIDER_EV / REV_FY26_CONS:.1f}x", f"Consensus ${REV_FY26_CONS:.1f}M; FY2027E ${REV_FY27_CONS:.1f}M after U.S. Ultomiris royalties end", "2026-10-27"),
    ("P/FCF, EV/FCF, EV/EBITDA", "n/m", f"TTM FCF ${FCF_TTM:.1f}M; OCF ${OCF_TTM:.1f}M; EBITDA −$205.9M", "2026-06-30"),
    ("P/B", f"{PRICE / BVPS:.2f}x", f"BVPS ${BVPS:.2f}; book is mostly cash and securities", "2026-06-30"),
    ("Net cash per share", f"${net_cash_ps:.2f}", f"Pipeline/platform premium the market pays: ${premium_ps:.2f}/share (~{premium_ps / PRICE:.0%} of price)", "2026-06-30"),
    ("Annualized cash burn", "~$200M", f"TTM FCF ${FCF_TTM:.0f}M; FY2026E FCF −$177M consensus; Alexion settlement adds $105M in two installments", "2026-10-27"),
    ("Short interest", "22.0% of shares / 30.2% of float", "11.3 days to cover; a crowded short into the XmAb819 data", "2026-10-28"),
    ("Analyst target (12M)", f"${ANALYST_PT:.2f} avg ($21-$70)", f"{ANALYST_PT / PRICE - 1:.0%} upside; 14 analysts (9 Strong Buy, 3 Buy, 2 Hold)", "2026-10-27"),
    ("Scenario targets (bear / base / bull)", f"${scen['Bear']['target']:.2f} / ${scen['Base']['target']:.2f} / ${scen['Bull']['target']:.2f}", "Sum-of-parts per share after scenario dilution", "2026-10-29"),
    ("Probability-weighted FV", f"${weighted_fv:.2f}", f"25% bear / 50% base / 25% bull; {weighted_fv / PRICE - 1:+.0%} vs $29.99", "2026-10-29"),
]
for r, row in enumerate(metrics, 11):
    for c, value in enumerate(row, 1):
        hl = r in (22, 23)
        put(ws, r, c, value, bold=hl, fill=GREEN if hl else None, wrap=True)
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 30
ws.column_dimensions["C"].width = 80
ws.column_dimensions["D"].width = 14
ws.freeze_panes = "A11"

ws = wb.create_sheet("WACC")
title(ws, "XNCR Weighted Average Cost of Capital", 4)
headers(ws, 3, ["Component", "Value", "Calculation / Source", "Comment"])
wacc_rows = [
    ("Risk-free rate", rf, "CNBC US10Y / Trading Economics", "10-year Treasury 5.02% on Oct. 28, 2026"),
    ("Equity risk premium", erp, "Model assumption", "Standard US ERP"),
    ("Levered beta", beta, "StockAnalysis statistics (5Y)", "Low beta understates binary clinical risk"),
    ("Cost of equity (CAPM)", ke, "Rf + beta × ERP", "Mechanical CAPM"),
    ("Pre-tax cost of debt", cost_debt, "OMERS royalty liability: $12.0M non-cash interest on ~$112M avg balance in H1 2026 (~21% annualized); leases ~6%", "Blended judgment; the royalty liability is non-recourse"),
    ("Tax rate", tax, "Large NOL position", "No tax shield"),
    ("Market capitalization", MARKET_CAP, "StockAnalysis", "$ millions"),
    ("Total debt", TOTAL_DEBT, "Royalty liability $104.9M + leases $65.2M", "$ millions"),
    ("Equity weight", equity_weight, "MC / (MC + debt)", ""),
    ("Debt weight", debt_weight, "Debt / (MC + debt)", ""),
    ("Computed WACC", wacc, "E/V × Ke + D/V × Kd × (1 − t)", "Provider estimate 9.42%"),
    ("Selected rNPV discount rate", rnpv_rate, "Judgmental clinical-stage rate", "Embedded in pipeline NPV/peak-sales multiples on Scenarios"),
]
for r, row in enumerate(wacc_rows, 4):
    for c, value in enumerate(row, 1):
        fmt = None
        if c == 2:
            fmt = "#,##0.0" if r in (10, 11) else ("0.00" if r == 6 else "0.00%")
        hl = r in (14, 15)
        put(ws, r, c, value, bold=hl, fill=GREEN if hl else None, fmt=fmt, wrap=True)
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 14
ws.column_dimensions["C"].width = 70
ws.column_dimensions["D"].width = 52

ws = wb.create_sheet("Scenarios")
title(ws, "XNCR Sum-of-Parts Scenario Analysis (Cash NAV + Risk-Adjusted Pipeline)", 6)
ws.merge_cells("A2:F2")
ws["A2"] = (
    "Primary framework: net cash + PV of the second Alexion settlement installment + royalty/milestone portfolio NPV "
    "+ XmAb819 rNPV (peak sales × probability of success × NPV/peak multiple at a 12% rate) + XmAb942 and other pipeline "
    "− PV of corporate overhead, divided by scenario-diluted shares. Revenue CAGR / FCF-multiple rows are omitted because "
    "Xencor has no product revenue and negative FCF by design."
)
ws["A2"].alignment = Alignment(wrap_text=True)
ws.row_dimensions[2].height = 60
headers(ws, 4, ["Metric", "Bear", "Base", "Bull", "Units / Formula", "Interpretation"])
B, M, U = scen["Bear"], scen["Base"], scen["Bull"]
scenario_rows = [
    ("Net cash (Jun. 30, 2026)", B["net_cash"], M["net_cash"], U["net_cash"], "$mm", "Cash & investments $565M − royalty liability − leases"),
    ("PV of 2nd Alexion installment", B["alexion"], M["alexion"], U["alexion"], "$mm", "$52.5M due after Jul. 2027; discounted at 12%"),
    ("Royalty & milestone portfolio NPV", B["royalty"], M["royalty"], U["royalty"], "$mm", "Ex-U.S. Ultomiris above OMERS caps, Monjuvi (cap reached), obexelimab royalties + up to $450M milestones, others"),
    ("XmAb819 peak sales (ccRCC + ENPP3+)", B["peak819"], M["peak819"], U["peak819"], "$mm", "Bear: niche late-line; base: 2L/3L ccRCC standard; bull: earlier lines and other ENPP3+ tumors"),
    ("XmAb819 probability of success", B["pos819"], M["pos819"], U["pos819"], "%", "Phase 1 cORR 45.5% (n=11) at 400 µg/kg; registrational study starts by YE2027"),
    ("XmAb819 NPV / peak-sales multiple", B["mult819"], M["mult819"], U["mult819"], "x", "Wholly owned; net of remaining development and launch cost at 12%"),
    ("XmAb819 rNPV", B["v819"], M["v819"], U["v819"], "$mm", "Peak × PoS × multiple"),
    ("XmAb942 (Xtend TL1A) value", B["v942"], M["v942"], U["v942"], "$mm", "XENITH-UC Phase 2b primary endpoint ~2H27; crowded TL1A class"),
    ("Other pipeline & platform", B["other"], M["other"], U["other"], "$mm", "XmAb412, XmAb541/808, XmAb657, plamotamab, XmAb discovery licensing"),
    ("PV of corporate overhead", B["corp"], M["corp"], U["corp"], "$mm", "G&A ~$65M/yr not allocated to programs"),
    ("Equity value", B["equity"], M["equity"], U["equity"], "$mm", "Sum of parts"),
    ("Dilution factor", B["dil"], M["dil"], U["dil"], "x", "Bear: raise at low prices; bull: data-driven raise or partnering funds Phase 3"),
    ("Diluted shares", B["shares"], M["shares"], U["shares"], "mm", "74.35M × dilution factor"),
    ("Pipeline value per share", B["pipeline_ps"], M["pipeline_ps"], U["pipeline_ps"], "$ / share", "XmAb819 + XmAb942 + other"),
    ("Target price", B["target"], M["target"], U["target"], "$ / share", "Equity value / diluted shares"),
    ("Upside / (downside)", B["target"] / PRICE - 1, M["target"] / PRICE - 1, U["target"] / PRICE - 1, "%", "Versus $29.99 close"),
    ("Probability", B["w"], M["w"], U["w"], "%", "Weights sum to 100%"),
    ("Weighted value / share", B["target"] * B["w"], M["target"] * M["w"], U["target"] * U["w"], "$ / share", "Target × probability"),
    ("Probability-weighted FV", None, weighted_fv, None, "$ / share", "Sum of weighted values"),
    ("Upside from current", None, weighted_fv / PRICE - 1, None, "%", f"Analyst average ${ANALYST_PT:.2f}"),
    ("Cross-check: net cash per share", None, net_cash_ps, None, "$ / share", "Liquidation-style floor before burn and dilution"),
    ("Cross-check: implied pipeline premium", None, premium_ps, None, "$ / share", "Price − net cash per share"),
]
for r, row in enumerate(scenario_rows, 5):
    for c, value in enumerate(row, 1):
        fmt = None
        if c in (2, 3, 4) and row[4] == "%":
            fmt = "0.0%"
        elif c in (2, 3, 4) and row[4] == "$ / share":
            fmt = "$#,##0.00"
        elif c in (2, 3, 4) and row[4] in ("$mm", "mm"):
            fmt = "#,##0.0"
        elif c in (2, 3, 4) and row[4] == "x":
            fmt = "0.00x"
        hl = r in (19, 23, 24)
        put(ws, r, c, value, bold=hl, fill=GREEN if hl else None, fmt=fmt, wrap=True)
ws.column_dimensions["A"].width = 40
for col in "BCD":
    ws.column_dimensions[col].width = 14
ws.column_dimensions["E"].width = 16
ws.column_dimensions["F"].width = 84
ws.freeze_panes = "A5"

urls = {
    "overview": "https://stockanalysis.com/stocks/xncr/",
    "statistics": "https://stockanalysis.com/stocks/xncr/statistics/",
    "income": "https://stockanalysis.com/stocks/xncr/financials/",
    "balance": "https://stockanalysis.com/stocks/xncr/financials/balance-sheet/",
    "cashflow": "https://stockanalysis.com/stocks/xncr/financials/cash-flow-statement/",
    "forecast": "https://stockanalysis.com/stocks/xncr/forecast/",
    "q2pr": "https://investors.xencor.com/news-releases/news-release-details/xencor-reports-second-quarter-2026-financial-results",
    "10q": "https://www.sec.gov/Archives/edgar/data/1326732/000132673226000056/xncr-20260630.htm",
    "esmo_pr": "https://investors.xencor.com/news-releases/news-release-details/xencor-presents-phase-1-dose-selection-data-xmab819-potential",
    "esmo_deck": "https://www.sec.gov/Archives/edgar/data/1326732/000121390026112279/ea030959601ex99-1.htm",
    "esmo_8k": "https://www.sec.gov/Archives/edgar/data/1326732/000121390026112279/ea030959601ex99-2.htm",
    "treasury": "https://www.cnbc.com/2026/10/28/us-treasurys-bonds-yields-federal-reserve.html",
    "fred": "https://fred.stlouisfed.org/series/dgs10",
}
ws = wb.create_sheet("Actuals Source Audit")
title(ws, "XNCR Actuals Source Audit", 5)
headers(ws, 3, ["Data point", "Value", "Source URL", "As of", "Notes"])
audit = [
    ("Close price", "$29.99 (after-hours $30.16)", urls["overview"], "2026-10-28", "52-wk $10.38-$33.60; +110% over 52 weeks; 50/200-DMA $25.65/$16.80"),
    ("Market cap / EV", "$2.23B / $1.83B", urls["statistics"], "2026-10-28", "EV = MC − $394.9M net cash"),
    ("Shares outstanding", "74.35M; +5.22% YoY", urls["statistics"], "2026-10-28", "Float 54.25M; institutions 88.3%; insiders 1.2%"),
    ("Beta", "0.90", urls["statistics"], "2026-10-28", "Five-year"),
    ("Revenue TTM / FY2025 / FY2024 / FY2023", "$104.97M / $125.58M / $110.49M / $174.62M", urls["income"], "2026-06-30", "FY2021 $275.1M peak on milestones"),
    ("Q2 2026 revenue mix", "Royalties $41.2M (incl. $27.6M Ultomiris catch-up); Zenas milestone $10.0M", urls["10q"], "2026-06-30", "Alexion $37.3M, Incyte $3.9M, Zenas $10.0M"),
    ("H1 2026 R&D / G&A", "$136.6M / $34.1M", urls["10q"], "2026-06-30", "XmAb942 $21.8M, XmAb819 $16.9M, XmAb412 $10.1M, XmAb541 $8.5M external R&D"),
    ("Operating income TTM", "−$216.08M", urls["income"], "2026-06-30", "StockAnalysis classifies R&D in cost of revenue (gross profit −$150.8M)"),
    ("Net income TTM / FY2025", "−$163.33M / −$91.92M", urls["income"], "2026-06-30", "TTM includes ~$63.9M gains on investments; FY2025 included ~$105M"),
    ("Equity-securities mark-to-market H1 2026", "−$33.9M", urls["10q"], "2026-06-30", "Q2 +$18.1M; holdings in partner stock create earnings noise"),
    ("OCF / capex / FCF TTM", "−$205.81M / −$2.47M / −$208.29M", urls["cashflow"], "2026-06-30", "SBC $39.0M"),
    ("Cash & investments", "$565.0M total; $486.4M cash + marketable debt securities", urls["q2pr"], "2026-06-30", "$610.8M at Dec. 31, 2025 (company definition)"),
    ("Marketable equity securities", "$78.6M", urls["10q"], "2026-06-30", "$112.5M at Dec. 31, 2025"),
    ("OMERS royalty liability", "$104.9M ($27.4M current)", urls["10q"], "2026-06-30", "Non-recourse; H1 non-cash interest $12.0M; Monjuvi cap ($29.3M) reached Q2 2026"),
    ("Ultomiris royalty sale terms", "OMERS gets up to $35M/yr of royalties 2026-2028, $12M/yr from 2029", urls["10q"], "2023-11-03", "Upfront $192.5M in Nov. 2023"),
    ("Lease liabilities", "$65.2M", urls["10q"], "2026-06-30", "Pasadena/Monrovia facilities"),
    ("Total debt (provider) / net cash", "$170.1M / $394.9M", urls["balance"], "2026-06-30", "Debt = royalty liability + leases; no bank debt"),
    ("Equity / BVPS", "$503.2M / $6.77", urls["balance"], "2026-06-30", "Accumulated deficit −$946.6M"),
    ("Alexion settlement", "$105.0M in two $52.5M installments (Aug. 2026, ~Aug. 2027)", urls["10q"], "2026-07-29", "Ends U.S. Ultomiris royalties; ex-U.S. royalties continue"),
    ("Zenas obexelimab", "BLA May 2026 ($10M milestone); Japan MAA Jul. 2026 ($2.5M, Q3)", urls["10q"], "2026-07-30", "Up to $450M milestones; mid-single to mid-teen royalties"),
    ("Cash guidance", "YE2026 $420-$440M; runway through 2028", urls["q2pr"], "2026-08-05", "Cash + marketable debt securities"),
    ("XmAb819 Phase 1", "119 pts, 11 dose levels; ORR 30% / 44% / 45.5% at 100 / 200 / 400 µg/kg", urls["esmo_pr"], "2026-09-01", "ESMO 2026 oral; data cut Sep. 1, 2026"),
    ("XmAb819 400 µg/kg (Oct. 1 cut)", "cORR 45.5% (5/11), uORR 55%, 6-mo PFS 82%, >70% on treatment", urls["esmo_pr"], "2026-10-01", "Proposed IV RP3D; registrational development by end of 2027"),
    ("XmAb942 XENITH-UC", "Phase 2b primary endpoint (12-week induction) ~2H27", urls["q2pr"], "2026-08-05", "Enrollment on track"),
    ("Consensus FY2026", "Revenue $148.84M; EPS −$3.07 (range −$1.82 to −$4.47)", urls["forecast"], "2026-10-27", "13 analysts; non-GAAP provider figures; revenue range $62.0M-$199.8M"),
    ("Consensus FY2027", "Revenue $121.03M; EPS −$3.54", urls["forecast"], "2026-10-27", "Headline only; detail gated"),
    ("Analyst targets", "$40.25 avg / $41.50 median / $21 low / $70 high", urls["forecast"], "2026-10-27", "RBC $20→$43, Wedbush $30→$45, Wells $30→$44, Barclays $35→$45, Truist $29→$40"),
    ("Short interest", "16.37M; 22.0% of shares; 30.2% of float; 11.3 days", urls["statistics"], "2026-10-28", "Previous month 15.75M"),
    ("Next earnings", "Nov. 9, 2026 (estimated, Q3 2026)", urls["statistics"], "2026-10-28", "Not yet confirmed by company"),
    ("US 10Y Treasury", "5.02%", urls["treasury"], "2026-10-28", "FRED DGS10 5.00% on Oct. 26"),
]
for r, row in enumerate(audit, 4):
    for c, value in enumerate(row, 1):
        put(ws, r, c, value, wrap=True)
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 50
ws.column_dimensions["C"].width = 70
ws.column_dimensions["D"].width = 14
ws.column_dimensions["E"].width = 60
ws.freeze_panes = "A4"

ws = wb.create_sheet("Questions")
title(ws, "XNCR Open Diligence Questions", 4)
headers(ws, 3, ["#", "Question", "Why it matters", "Evidence needed"])
questions = [
    (1, "How durable are XmAb819 responses at 400 µg/kg once the n=11 efficacy cohort matures?", "The 45.5% cORR rests on 5 confirmed responses; a few conversions or progressions move the rate by 10 points.", "Median DoR and PFS at the next update; larger 400 µg/kg cohort"),
    (2, "What exactly is the registrational design and comparator, and will FDA accept accelerated approval on sub-group data?", "Timing to revenue and the cost of Phase 3 drive most of XmAb819's NPV.", "End-of-Phase 1/2 FDA meeting minutes; trial registration"),
    (3, "How manageable is cytokine release syndrome, rash and transaminase elevation with the weekly priming regimen in community oncology?", "T-cell engagers in solid tumors have repeatedly failed on tolerability or step-up dosing logistics.", "Grade 3+ CRS rate, hospitalization requirements, SC formulation plans"),
    (4, "What share of ccRCC patients express ENPP3 at levels that respond, and is a companion diagnostic needed?", "Addressable market and label breadth depend on expression thresholds.", "Biomarker response analysis; IHC assay plans"),
    (5, "Why did U.S. Ultomiris royalties end via a $105M settlement, and what ex-U.S. royalties remain after OMERS takes up to $35M a year through 2028?", "Royalty income was the main non-dilutive funding source; FY2027 revenue consensus falls to $121M.", "Ex-U.S. Ultomiris sales trajectory; settlement terms"),
    (6, "Is the OMERS royalty liability ($104.9M) fully non-recourse, and when does it stop absorbing royalty cash?", "It is counted as debt by data providers and accretes at a ~21% effective rate.", "Note 6 detail; payment schedule"),
    (7, "Does cash runway 'through 2028' include a registrational XmAb819 study and the XmAb942 Phase 3 decision?", "If not, a raise is likely after the stock's ~3x move, and timing determines dilution.", "Q3 2026 guidance update; ATM usage"),
    (8, "Will Xencor partner XmAb819 or XmAb942, and on what terms?", "A partnership could fund Phase 3 non-dilutively but cede economics.", "Management commentary; BD announcements"),
    (9, "How differentiated is XmAb942's Xtend half-life against Merck, Roche and Sanofi TL1A antibodies already in Phase 3?", "Less-frequent dosing may not be enough in a crowded class; XmAb942 consumed the most external R&D in H1 2026.", "XENITH-UC efficacy vs competitor Phase 2 data; dosing interval"),
    (10, "What are the marketable equity securities ($78.6M) and is the company a forced holder?", "Mark-to-market swung −$33.9M in H1 2026, making net income noisy.", "10-Q fair-value table by issuer"),
    (11, "What is obexelimab's probability of approval in IgG4-RD and what royalty tier applies?", "Approval would convert milestones into a recurring royalty stream.", "FDA PDUFA date; Zenas launch plan"),
    (12, "Why did revenue fall from $275M (FY2021) to ~$105M TTM, and how much of the remaining base is recurring?", "Milestones are lumpy; investors may extrapolate the Q2 2026 royalty catch-up.", "Royalty vs milestone split by year"),
    (13, "What explains short interest of 30% of float with 11 days to cover?", "Shorts may be positioned for tolerability issues, Phase 3 failure or an equity raise.", "Short borrow cost; options positioning"),
    (14, "How will management allocate R&D across 819, 942, 412, 541/808, 657 and plamotamab if 819 needs a large Phase 3?", "Breadth spreads cash thin; prioritization signals discipline.", "2027 guidance; program terminations"),
    (15, "Is stock-based compensation ($39M TTM) and 5% annual share growth likely to continue?", "Ongoing dilution reduces per-share value of any success.", "Proxy statement; share count trend"),
]
for r, row in enumerate(questions, 4):
    for c, value in enumerate(row, 1):
        put(ws, r, c, value, wrap=True)
ws.column_dimensions["A"].width = 7
ws.column_dimensions["B"].width = 72
ws.column_dimensions["C"].width = 62
ws.column_dimensions["D"].width = 45

ws = wb.create_sheet("Sources")
title(ws, "XNCR Sources", 4)
headers(ws, 3, ["#", "Source", "URL", "Use"])
source_rows = [
    (1, "StockAnalysis overview", urls["overview"], "Price, news, earnings date"),
    (2, "StockAnalysis statistics", urls["statistics"], "Market value, ratios, shares, beta, short interest"),
    (3, "StockAnalysis income statement", urls["income"], "Revenue, operating income, net income history"),
    (4, "StockAnalysis balance sheet", urls["balance"], "Cash, investments, royalty liability, leases, equity"),
    (5, "StockAnalysis cash flow", urls["cashflow"], "OCF, capex, FCF, SBC"),
    (6, "StockAnalysis forecast", urls["forecast"], "Consensus revenue/EPS, targets and ratings"),
    (7, "Xencor Q2 2026 earnings release", urls["q2pr"], "Cash guidance, pipeline milestones"),
    (8, "Xencor Q2 2026 Form 10-Q", urls["10q"], "Revenue by partner, R&D by program, OMERS, Alexion settlement, Zenas"),
    (9, "Xencor XmAb819 ESMO press release (Oct. 25, 2026)", urls["esmo_pr"], "Phase 1 dose-selection data"),
    (10, "Xencor XmAb819 investor deck (8-K Ex. 99.1)", urls["esmo_deck"], "Pipeline milestones, safety data, development plan"),
    (11, "Xencor XmAb819 press release (8-K Ex. 99.2)", urls["esmo_8k"], "Registrational study framing"),
    (12, "CNBC Treasury yields Oct. 28, 2026", urls["treasury"], "WACC risk-free rate"),
    (13, "FRED DGS10", urls["fred"], "Risk-free rate cross-check"),
]
for r, row in enumerate(source_rows, 4):
    for c, value in enumerate(row, 1):
        put(ws, r, c, value, wrap=True)
ws.column_dimensions["A"].width = 7
ws.column_dimensions["B"].width = 46
ws.column_dimensions["C"].width = 100
ws.column_dimensions["D"].width = 60

for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = False

wb.save(OUT)
check = load_workbook(OUT, data_only=False)
expected = ["Valuation", "WACC", "Scenarios", "Actuals Source Audit", "Questions", "Sources"]
assert check.sheetnames == expected, check.sheetnames
assert abs(check["Scenarios"]["C23"].value - weighted_fv) < 1e-6
assert scen["Bear"]["target"] < PRICE < scen["Bull"]["target"]
print(f"WACC: {wacc:.2%} (Ke {ke:.2%}, Kd {cost_debt:.2%}); rNPV rate {rnpv_rate:.0%}")
print(f"Net cash ${NET_CASH:.1f}M = ${net_cash_ps:.2f}/sh; pipeline premium ${premium_ps:.2f}/sh")
for k, s in scen.items():
    print(f"{k}: 819 rNPV ${s['v819']:,.0f}M, equity ${s['equity']:,.0f}M, shares {s['shares']:.1f}M, target ${s['target']:.2f} ({s['target'] / PRICE - 1:+.1%})")
print(f"Probability-weighted FV: ${weighted_fv:.2f} ({weighted_fv / PRICE - 1:+.1%}); analyst avg ${ANALYST_PT}")
print(f"Created {OUT}")
