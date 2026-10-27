#!/usr/bin/env python3
"""Build the 2026-10-27 Nebius Group valuation workbook."""

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


OUT = Path(__file__).with_name("2026-10-27 Nebius Group Model.xlsx")
PRICE = 153.27  # Oct. 26, 2026 close
SHARES = 274.10  # millions, Class A + Class B
MARKET_CAP = 42010.0  # $ millions
CASH = 8042.0  # Jun. 30, 2026
DEBT = 10195.0  # includes $1,625M leases; excludes Jul. $775M ABS and Oct. $3.5B Citi bridge
NET_DEBT = DEBT - CASH
PROVIDER_EV = 44160.0
UNEARNED = 1068.0 + 4996.0  # customer prepayments not yet recognized
REV_TTM = 1355.0
REV_FY25 = 529.8
REV_FY26_CONS = 3350.0
REV_FY27_CONS = 12240.0
EPS_FY26_CONS = -2.10  # non-GAAP, S&P Global
EPS_FY27_CONS = -3.60
ADJ_EBITDA_Q2 = 236.2
ADJ_EBITDA_H1 = 365.7
ARR_Q2 = 3000.0
OCF_TTM = 5258.0
CAPEX_TTM = 11142.0
FCF_TTM = OCF_TTM - CAPEX_TTM
OP_INC_TTM = -507.3
NI_TTM = 42.4
BVPS = 37.72
ANALYST_PT = 281.0
TAX = 0.21
HORIZON = 5

NAVY = "17365D"
BLUE = "D9EAF7"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
WHITE = "FFFFFF"
thin = Side(style="thin", color="B7B7B7")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def put(ws, row, col, value, *, bold=False, fill=None, fmt=None, wrap=False):
    cell = ws.cell(row=row, column=col, value=value)
    cell.border = border
    cell.font = Font(bold=bold, color=WHITE if fill == NAVY else "000000")
    cell.alignment = Alignment(vertical="top", wrap_text=wrap)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if fmt:
        cell.number_format = fmt
    return cell


def title(ws, text, end_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    cell = ws.cell(1, 1, text)
    cell.font = Font(size=15, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 24


def headers(ws, row, values):
    for col, value in enumerate(values, 1):
        put(ws, row, col, value, bold=True, fill=BLUE, wrap=True)


# WACC
rf = 0.04978
erp = 0.05
beta = 1.46
ke = rf + beta * erp
cost_debt = 0.065  # blend: 1.00%/2.75% converts, SOFR+2.50% ABS, SOFR+2.75-3.50% Citi bridge
equity_weight = MARKET_CAP / (MARKET_CAP + DEBT)
debt_weight = 1 - equity_weight
wacc = equity_weight * ke + debt_weight * cost_debt * (1 - TAX)

# Scenario inputs: FY2031 adjusted EBITDA x exit EV/EBITDA, less FY2031 net debt, over FY2031 diluted shares.
scen = {
    "Bear": {"rev": 14000.0, "margin": 0.45, "mult": 6.5, "net_debt": 18000.0, "shares": 335.0, "w": 0.30},
    "Base": {"rev": 30000.0, "margin": 0.55, "mult": 8.5, "net_debt": 35000.0, "shares": 330.0, "w": 0.50},
    "Bull": {"rev": 40000.0, "margin": 0.57, "mult": 9.5, "net_debt": 45000.0, "shares": 315.0, "w": 0.20},
}
for s in scen.values():
    s["cagr"] = (s["rev"] / REV_FY26_CONS) ** (1 / HORIZON) - 1
    s["ebitda"] = s["rev"] * s["margin"]
    s["ev"] = s["ebitda"] * s["mult"]
    s["equity"] = s["ev"] - s["net_debt"]
    s["target"] = s["equity"] / s["shares"]
    s["dilution"] = s["shares"] / SHARES - 1
    s["pv"] = s["target"] / (1 + ke) ** HORIZON
weighted_fv = sum(s["target"] * s["w"] for s in scen.values())
weighted_pv = weighted_fv / (1 + ke) ** HORIZON

wb = Workbook()
ws = wb.active
ws.title = "Valuation"
title(ws, "Nebius Group N.V. (NASDAQ: NBIS) — Valuation Summary", 4)
summary = [
    ("Company", "Nebius Group N.V.", "Model date", "2026-10-27"),
    ("Quote", "$153.27 at Oct. 26, 2026 close", "Shares outstanding", "274.10M (Class A 238.40M + Class B)"),
    ("Market capitalization", "$42.01B", "Enterprise value", f"${PROVIDER_EV / 1000:.2f}B (MC + $10.20B debt incl. leases − $8.04B cash)"),
    ("Net debt (Jun. 30, 2026)", f"${NET_DEBT / 1000:.2f}B", "Post-quarter financing", "$775M ABS (Jul.) and $3.5B Citi 364-day bridge (Oct.) not yet in reported debt"),
    ("Primary lens", "FY2031 adjusted EBITDA × exit EV/EBITDA, less projected net debt; EV/ARR and EV/forward sales cross-checks", "Stance", "Watch / Positive bias"),
]
for r, row in enumerate(summary, 3):
    for c, value in enumerate(row, 1):
        put(ws, r, c, value, bold=c in (1, 3), fill=YELLOW if c in (1, 3) else None, wrap=True)

headers(ws, 10, ["Metric", "Current / Forward", "Interpretation", "Source date"])
metrics = [
    ("Trailing P/E (GAAP)", f"{PRICE / 0.16:.0f}x", "$0.16 TTM EPS includes $780.6M Q1 2026 investment-sale gain; operating loss TTM −$507.3M. Not meaningful", "2026-06-30"),
    ("FY2026E / FY2027E EPS (consensus)", f"${EPS_FY26_CONS:.2f} / ${EPS_FY27_CONS:.2f}", "Non-GAAP losses widen as D&A and interest ramp ahead of revenue; P/E unusable", "2026-10-22"),
    ("EV / TTM revenue", f"{PROVIDER_EV / REV_TTM:.1f}x", f"TTM revenue ${REV_TTM:,.0f}M (+507%)", "2026-06-30"),
    ("EV / FY2026E revenue", f"{PROVIDER_EV / REV_FY26_CONS:.1f}x", "$3.35B consensus (21 analysts); guidance $3.0-$3.4B", "2026-10-22"),
    ("EV / FY2027E revenue", f"{PROVIDER_EV / REV_FY27_CONS:.1f}x", "$12.24B consensus headline (+266%); ranges gated", "2026-10-22"),
    ("EV / Q2 ARR", f"{PROVIDER_EV / ARR_Q2:.1f}x", "$3.0B ARR at June; YE2026 guidance $7-$9B implies ~5.5x at midpoint", "2026-06-30"),
    ("EV / annualized Q2 adj. EBITDA", f"{PROVIDER_EV / (ADJ_EBITDA_Q2 * 4):.1f}x", f"Q2 adj. EBITDA ${ADJ_EBITDA_Q2}M (41% margin); AI cloud 49.7%", "2026-06-30"),
    ("EV / FY2026E adj. EBITDA", f"{PROVIDER_EV / (REV_FY26_CONS * 0.40):.1f}x", "Guidance ~40% margin on consensus revenue", "2026-10-22"),
    ("P/B", f"{PRICE / BVPS:.2f}x", f"BVPS ${BVPS}; TBVPS $35.21 after $2.8B ATM at $223.60", "2026-06-30"),
    ("FCF yield (TTM)", f"{FCF_TTM / MARKET_CAP:.1%}", f"OCF ${OCF_TTM:,.0f}M (incl. ~$5.96B prepayments) − capex ${CAPEX_TTM:,.0f}M", "2026-06-30"),
    ("Customer prepayments on balance sheet", f"${UNEARNED / 1000:.2f}B", "Current $1.07B + long-term $5.00B unearned revenue; future revenue already collected in cash", "2026-06-30"),
    ("Short interest", "16.31% of shares / 19.54% of float", "44.70M shares; 2.98 days to cover", "2026-10-26"),
    ("Analyst target (12M)", f"${ANALYST_PT:.0f} avg ($84-$410)", f"{ANALYST_PT / PRICE - 1:.0%} upside; 22 analysts, consensus Buy", "2026-10-22"),
    ("Probability-weighted FV (FY2031)", f"${weighted_fv:.2f}", f"30% bear / 50% base / 20% bull; {weighted_fv / PRICE - 1:.0%} undiscounted", "2026-10-27"),
    ("PV of weighted FV at cost of equity", f"${weighted_pv:.2f}", f"Discounted 5 years at {ke:.2%}; {weighted_pv / PRICE - 1:.0%} vs current — modest excess return for high risk", "2026-10-27"),
]
for r, row in enumerate(metrics, 11):
    for c, value in enumerate(row, 1):
        hl = r in (24, 25)
        put(ws, r, c, value, bold=hl, fill=GREEN if hl else None, wrap=True)
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 30
ws.column_dimensions["C"].width = 78
ws.column_dimensions["D"].width = 14
ws.freeze_panes = "A11"

ws = wb.create_sheet("WACC")
title(ws, "NBIS Weighted Average Cost of Capital", 4)
headers(ws, 3, ["Component", "Value", "Calculation / Source", "Comment"])
wacc_rows = [
    ("Risk-free rate", rf, "CNBC US10Y", "Oct. 26-27, 2026"),
    ("Equity risk premium", erp, "Model assumption", "Standard US ERP"),
    ("Levered beta", beta, "StockAnalysis statistics", "Five-year beta; trading history mixes pre-2024 Yandex era"),
    ("Cost of equity (CAPM)", ke, "Rf + beta × ERP", "Arguably understated for a levered, pre-profit neocloud"),
    ("Pre-tax cost of debt", cost_debt, "Blend of 1.00% 2030 and 2.75% 2032 converts, SOFR+2.50% ABS, SOFR+2.75-3.50% Citi bridge", "Assumption; marginal secured cost is higher than the convert coupons"),
    ("Tax rate", TAX, "US statutory", "Netherlands statutory 25.8%; TTM effective 19.4%"),
    ("Market capitalization", MARKET_CAP, "StockAnalysis", "$ millions"),
    ("Total debt", DEBT, "StockAnalysis balance sheet, Jun. 30, 2026", "$ millions; includes leases; excludes post-quarter $4.3B"),
    ("Equity weight", equity_weight, "MC / (MC + debt)", ""),
    ("Debt weight", debt_weight, "Debt / (MC + debt)", ""),
    ("Computed WACC", wacc, "E/V × Ke + D/V × Kd × (1 − t)", "Provider estimate 10.29%"),
]
for r, row in enumerate(wacc_rows, 4):
    for c, value in enumerate(row, 1):
        fmt = None
        if c == 2:
            fmt = "#,##0.0" if r in (10, 11) else ("0.00" if r == 6 else "0.00%")
        put(ws, r, c, value, bold=r == 14, fill=GREEN if r == 14 else None, fmt=fmt, wrap=True)
ws.column_dimensions["A"].width = 28
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 60
ws.column_dimensions["D"].width = 58

ws = wb.create_sheet("Scenarios")
title(ws, "NBIS Five-Year EV/EBITDA Scenario Analysis (FY2031)", 6)
ws.merge_cells("A2:F2")
ws["A2"] = (
    "Primary framework: FY2031 adjusted EBITDA × exit EV/EBITDA − projected FY2031 net debt, divided by projected "
    "diluted shares. FCF multiples are not used: TTM FCF is −$5.9B in a deliberate capex cycle ($20-25B 2026 guide), "
    "and GAAP/non-GAAP EPS stays negative through FY2027 consensus. Net debt and share count carry the funding cost of the build-out."
)
ws["A2"].alignment = Alignment(wrap_text=True)
ws.row_dimensions[2].height = 48
headers(ws, 4, ["Metric", "Bear", "Base", "Bull", "Units / Formula", "Interpretation"])
B, M, U = scen["Bear"], scen["Base"], scen["Bull"]
scenario_rows = [
    ("FY2031 revenue", B["rev"], M["rev"], U["rev"], "$mm", "Bear: GPU pricing resets and capacity slips; base: ~2 GW billing at $15M/MW; bull: >1 GW/yr deployment holds pricing"),
    ("Revenue CAGR (FY2026E-FY2031)", B["cagr"], M["cagr"], U["cagr"], "%", "From $3.35B FY2026 consensus; FY2027 consensus is $12.24B"),
    ("Adjusted EBITDA margin", B["margin"], M["margin"], U["margin"], "%", "Q2 2026 group 41%, AI cloud 49.7%; FY2026 guide ~40%"),
    ("FY2031 adjusted EBITDA", B["ebitda"], M["ebitda"], U["ebitda"], "$mm", "Revenue × margin"),
    ("Exit EV/EBITDA", B["mult"], M["mult"], U["mult"], "x", "GPU assets depreciate over five years; mature infra leasing trades 6-10x"),
    ("Implied enterprise value", B["ev"], M["ev"], U["ev"], "$mm", "EBITDA × multiple"),
    ("Less FY2031 net debt", B["net_debt"], M["net_debt"], U["net_debt"], "$mm", "From $2.15B today; cumulative negative FCF less equity raised"),
    ("Equity value", B["equity"], M["equity"], U["equity"], "$mm", "EV − net debt"),
    ("FY2031 diluted shares", B["shares"], M["shares"], U["shares"], "mm", "274.1M today; ATM, converts and SBC dilution"),
    ("Dilution vs today", B["dilution"], M["dilution"], U["dilution"], "%", "Bear issues more shares at lower prices"),
    ("Target price (FY2031)", B["target"], M["target"], U["target"], "$ / share", "Equity value / shares"),
    ("Upside / (downside)", B["target"] / PRICE - 1, M["target"] / PRICE - 1, U["target"] / PRICE - 1, "%", "Versus $153.27 close; undiscounted"),
    ("Probability", B["w"], M["w"], U["w"], "%", "Weights sum to 100%"),
    ("Weighted value / share", B["target"] * B["w"], M["target"] * M["w"], U["target"] * U["w"], "$ / share", "Target × probability"),
    ("Probability-weighted FV (FY2031)", None, weighted_fv, None, "$ / share", "Sum of weighted values"),
    ("Upside from current", None, weighted_fv / PRICE - 1, None, "%", "Undiscounted five-year"),
    ("PV of weighted FV at Ke", None, weighted_pv, None, "$ / share", f"Discounted {HORIZON} years at {ke:.2%}"),
    ("PV upside from current", None, weighted_pv / PRICE - 1, None, "%", "Risk-adjusted excess return"),
    ("Cross-check: PV of each target at Ke", B["pv"], M["pv"], U["pv"], "$ / share", "Bear and base PV bracket today's price"),
    ("Cross-check: EV / FY2031 revenue", B["ev"] / B["rev"], M["ev"] / M["rev"], U["ev"] / U["rev"], "x", f"Today {PROVIDER_EV / REV_FY27_CONS:.1f}x FY2027E revenue"),
]
for r, row in enumerate(scenario_rows, 5):
    for c, value in enumerate(row, 1):
        fmt = None
        if c in (2, 3, 4) and row[4] == "%":
            fmt = "0.0%"
        elif c in (2, 3, 4) and row[4] == "$ / share":
            fmt = "$#,##0.00"
        elif c in (2, 3, 4) and row[4] in ("$mm", "mm"):
            fmt = "#,##0"
        elif c in (2, 3, 4) and row[4] == "x":
            fmt = "0.0x"
        hl = r in (19, 20, 21, 22)
        put(ws, r, c, value, bold=r == 15 or hl, fill=GREEN if hl else None, fmt=fmt, wrap=True)
ws.column_dimensions["A"].width = 38
for col in "BCD":
    ws.column_dimensions[col].width = 14
ws.column_dimensions["E"].width = 16
ws.column_dimensions["F"].width = 78
ws.freeze_panes = "A5"

urls = {
    "overview": "https://stockanalysis.com/stocks/nbis/",
    "income": "https://stockanalysis.com/stocks/nbis/financials/",
    "quarterly": "https://stockanalysis.com/stocks/nbis/financials/?p=quarterly",
    "balance": "https://stockanalysis.com/stocks/nbis/financials/balance-sheet/?p=quarterly",
    "cashflow": "https://stockanalysis.com/stocks/nbis/financials/cash-flow-statement/",
    "statistics": "https://stockanalysis.com/stocks/nbis/statistics/",
    "forecast": "https://stockanalysis.com/stocks/nbis/forecast/",
    "q2letter": "https://www.sec.gov/Archives/edgar/data/1513845/000110465926094568/tm2622968d1_ex99-2.htm",
    "q2pr": "https://www.businesswire.com/news/home/20260812963359/en/Nebius-reports-second-quarter-2026-financial-results",
    "q2call": "https://www.tradingkey.com/news/transcripts/262106466-tradingkey",
    "q1call": "https://finance.yahoo.com/quote/NBIS/earnings/NBIS-Q1-2026-earnings_call-555850.html/",
    "q4letter": "https://www.sec.gov/Archives/edgar/data/1513845/000110465926013946/tm266173d1_ex99-2.htm",
    "citi": "https://www.sec.gov/Archives/edgar/data/1513845/000110465926117978/tm2628078d1_6k.htm",
    "citipr": "https://nebius.com/newsroom/nebius-secures-3-5-billion-financing-facility-with-citi",
    "msft": "https://www.sec.gov/Archives/edgar/data/1513845/000110465925088312/tm2525580d1_6k.htm",
    "converts": "https://nebius.com/newsroom/nebius-group-announces-pricing-of-upsized-private-offering-of-2-75-billion-of-convertible-senior-notes",
    "meta": "https://www.cnbc.com/2026/03/16/meta-nebius-ai-infrastructure.html",
    "ceo": "https://finance.yahoo.com/markets/stocks/articles/nebius-ceo-admits-neocloud-financing-041453980.html",
    "treasury": "https://www.cnbc.com/quotes/US10Y",
}
ws = wb.create_sheet("Actuals Source Audit")
title(ws, "NBIS Actuals Source Audit", 5)
headers(ws, 3, ["Data point", "Value", "Source URL", "As of", "Notes"])
audit = [
    ("Close price", "$153.27 (52-wk $73.52-$299.86)", urls["overview"], "2026-10-26", "After-hours $155.89; ~33% down in October per Yahoo/Benzinga"),
    ("Market cap / EV", "$42.01B / $44.16B", urls["statistics"], "2026-10-26", "EV = MC + $10.20B debt − $8.04B cash"),
    ("Shares outstanding", "274.10M; +22.12% YoY", urls["statistics"], "2026-10-26", "Class A 238.40M; insiders 15.21%; float 228.69M"),
    ("Beta", "1.46", urls["statistics"], "2026-10-26", "Five-year"),
    ("Revenue TTM / FY2025 / FY2024", "$1,355M / $529.8M / $91.5M", urls["income"], "2026-06-30", "FY2021 $4.75B was pre-divestiture Yandex"),
    ("Quarterly revenue ($M)", "146.1 / 227.7 / 399.0 / 582.3", urls["quarterly"], "3Q25-2Q26", "AI cloud 131.1 / 214.1 / 389.7 / 574.9"),
    ("Gross profit / margin TTM", "$1,006M / 74.3%", urls["income"], "2026-06-30", "Provider gross profit excludes D&A"),
    ("Operating income TTM / FY2025", "−$507.3M / −$611.7M", urls["income"], "2026-06-30", "Q2 2026 −$1.3M incl. $115.9M Eigen AI SBC/post-combination"),
    ("Net income TTM / FY2025", "$42.4M / $82.5M", urls["income"], "2026-06-30", "Q1 2026 $780.6M and Q2 2025 $597.4M investment-sale gains"),
    ("D&A TTM / Q2 2026", "$863.8M / $294.3M (CF basis)", urls["cashflow"], "2026-06-30", "Server useful life raised from 4 to 5 years in 2026"),
    ("SBC TTM", "$188.9M", urls["cashflow"], "TTM 2026-06-30", "Q2 2026 $102.5M"),
    ("Adjusted EBITDA Q2 / H1 2026", "$236.2M (41%) / $365.7M (37%)", urls["q2letter"], "2026-06-30", "AI cloud $285.7M, 49.7%"),
    ("Adjusted net loss Q2 2026", "−$33.2M", urls["q2pr"], "2026-06-30", "GAAP net loss from continuing ops −$190.4M"),
    ("ARR", "$3.0B Jun.; $1.9B Mar.; $1.25B Dec.", urls["q2letter"], "2026-06-30", "Last month AI cloud revenue × 12"),
    ("FY2026 guidance (reaffirmed)", "Revenue $3.0-$3.4B; ARR $7-$9B; adj. EBITDA ~40%; capex $20-$25B", urls["q2call"], "2026-08-12", "Capex raised from $16-$20B in May"),
    ("Contracted power", "5 GW targeted for YE2026", urls["q2letter"], "2026-08-12", ">1 GW/yr deployment planned from 2027"),
    ("Customer commitments / prepayments", ">$40B commitments; >$9B prepayments expected in 2026", urls["q2letter"], "2026-08-12", "70% of Q2 deals included prepayments covering 50-60% of capex"),
    ("Microsoft contract", "~$17.4B through 2031 (up to $19.4B)", urls["msft"], "2025-09-08", "Vineland, NJ; all tranches delivered by Aug. 2026"),
    ("Meta contract", "Up to $27B over five years ($12B dedicated)", urls["meta"], "2026-03-16", "Second Meta build on track for early 2027"),
    ("Cash / debt", "$8,042M / $10,195M", urls["balance"], "2026-06-30", "LT debt $8,499M; leases $1,625M; restricted cash $830M"),
    ("Unearned revenue", "$1,068M current + $4,996M long-term", urls["balance"], "2026-06-30", "Customer prepayments"),
    ("PP&E / total assets", "$14,900M / $27,962M", urls["balance"], "2026-06-30", "PP&E was $891.5M at Dec. 2024"),
    ("Equity / BVPS", "$10,340M / $37.72", urls["balance"], "2026-06-30", "APIC +$3.0B in Q2 from ATM"),
    ("OCF / capex TTM", "$5,258M / −$11,142M", urls["cashflow"], "TTM 2026-06-30", "Unearned revenue change +$5,958M inside OCF"),
    ("FCF TTM / FY2025", "−$5,884M / −$3,681M", urls["cashflow"], "2026-06-30", "Q2 2026 −$3,411M"),
    ("Equity issued TTM", "$4,161M", urls["cashflow"], "TTM 2026-06-30", "ATM 12.7M shares at $223.60 = ~$2.8B gross"),
    ("Debt issued TTM", "$7,500M", urls["cashflow"], "TTM 2026-06-30", "Q3 2025 $3,163M (incl. $2.75B converts); Q1 2026 $4,338M"),
    ("Convertible notes", "$1.375B 1.00% 2030 + $1.375B 2.75% 2032", urls["converts"], "2025-09-10", "Plus later issuance; terms of Q1 2026 debt not reviewed"),
    ("Asset-backed financing", "~$775M at SOFR + 2.50%", urls["q2letter"], "2026-07", "Backed by GPUs and an IG customer contract"),
    ("Citi facility", "$3.5B senior secured, 364-day, SOFR + 2.75% rising to 3.50%", urls["citi"], "2026-10-17", "Bridge Facility Agreement"),
    ("Consensus FY2026 / FY2027", "Rev $3.35B / $12.24B; EPS −$2.10 / −$3.60", urls["forecast"], "2026-10-22", "Non-GAAP, S&P Global; 21 analysts FY2026; FY2026 FCF −$17.54B"),
    ("Analyst targets", "$281 avg; $302 median; $84-$410", urls["forecast"], "2026-10-22", "22 analysts: 13 SB, 2 B, 6 H, 1 SS"),
    ("Short interest", "44.70M; 16.31% of shares", urls["statistics"], "2026-10-26", "Prior month 46.78M"),
    ("Next earnings", "Nov. 10, 2026 (Q3 2026)", urls["overview"], "2026-10-26", ""),
    ("US 10Y Treasury", "4.978%", urls["treasury"], "2026-10-27", "CNBC quote page"),
]
for r, row in enumerate(audit, 4):
    for c, value in enumerate(row, 1):
        put(ws, r, c, value, wrap=True)
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 44
ws.column_dimensions["C"].width = 70
ws.column_dimensions["D"].width = 16
ws.column_dimensions["E"].width = 55
ws.freeze_panes = "A4"

ws = wb.create_sheet("Questions")
title(ws, "NBIS Open Diligence Questions", 4)
headers(ws, 3, ["#", "Question", "Why it matters", "Evidence needed"])
questions = [
    (1, "How will the $3.5B 364-day Citi bridge be refinanced, and at what cost and seniority?", "A bridge maturing in October 2027 makes 2027 refinancing the key solvency test while credit to neoclouds tightens.", "Term-out structure, covenants, Q3 call commentary"),
    (2, "Can H2 2026 exit at $7-$9B ARR from $3.0B in June?", "Requires capacity for Q2 deals to go live on time; slippage pushes revenue right while capex is already spent.", "Q3 ARR, MW live, delivery schedule"),
    (3, "What instrument made up the $4.34B of debt issued in Q1 2026, and what are its conversion terms?", "Convert dilution and coupons feed the share count and interest lines.", "20-F/6-K note disclosures"),
    (4, "What drove the $780.6M Q1 2026 and $597.4M Q2 2025 investment-sale gains, and how much ClickHouse/Toloka value remains?", "Gains made GAAP net income positive; remaining stakes are off-model optionality ($1.62B carried).", "Investment footnote, stake percentages"),
    (5, "How much of TTM OCF ($5.26B) is customer prepayments ($5.96B unearned revenue change)?", "Operating cash before prepayments is negative; prepaid revenue will be recognized without new cash.", "Prepayment schedule, revenue recognition timing"),
    (6, "What is the effect of extending server useful life from four to five years?", "Lowers D&A and flatters operating income; GPU economic life may be shorter as Rubin ships.", "Quantified D&A impact, impairment testing"),
    (7, "What share of revenue and backlog comes from Microsoft and Meta?", "Concentration in two hyperscalers that also build their own capacity.", "Customer concentration disclosure in 20-F"),
    (8, "Are $20-25M/MW contract yields and the 1y10m payback sustainable as Blackwell supply loosens?", "Pricing is the variable separating base and bear.", "Pricing on renewals, auction results, older-GPU pricing"),
    (9, "How much capacity is owned vs colocated vs asset-light partner, and what are lease obligations?", "Leases ($1.6B) and colocation commitments are debt-like and grow with the build.", "Commitments table"),
    (10, "How many ATM shares were sold after June 30 (12.3M remained), and at what price?", "Shares fell from ~$224 average ATM price to $153; further issuance is more dilutive.", "Q3 share count"),
    (11, "What is the plan for Avride funding and TripleTen?", "Non-core businesses consume cash and management attention.", "Segment adjusted EBITDA"),
    (12, "What governance protections apply given Class B supervoting shares held by the founder?", "Minority holders bear dilution decisions without control.", "20-F share class terms"),
    (13, "How much restricted cash ($830M) is pledged to secured facilities?", "Reduces true liquidity.", "Debt footnotes"),
    (14, "What portion of 2027 capacity is already sold and at what duration?", "Management says it could sell all 2027 capacity; mix of short vs long term sets revenue visibility.", "Backlog/RPO disclosure"),
    (15, "What happens to contract economics if a customer like an AI neolab fails to pay?", "Several Q2 landmark customers are venture-funded labs.", "Credit quality of counterparties, prepayment coverage"),
    (16, "What is the maintenance capex need once the build slows?", "Five-year GPU refresh cycles mean FCF may never reach EBITDA-like levels.", "Capex split by growth vs replacement"),
]
for r, row in enumerate(questions, 4):
    for c, value in enumerate(row, 1):
        put(ws, r, c, value, wrap=True)
ws.column_dimensions["A"].width = 7
ws.column_dimensions["B"].width = 72
ws.column_dimensions["C"].width = 62
ws.column_dimensions["D"].width = 45

ws = wb.create_sheet("Sources")
title(ws, "NBIS Sources", 4)
headers(ws, 3, ["#", "Source", "URL", "Use"])
source_rows = [
    (1, "StockAnalysis overview", urls["overview"], "Price, range, news, earnings date"),
    (2, "StockAnalysis statistics", urls["statistics"], "Market value, ratios, shares, beta, short interest"),
    (3, "StockAnalysis forecast", urls["forecast"], "Consensus revenue/EPS, targets and ratings"),
    (4, "StockAnalysis income statement", urls["income"], "Annual revenue, segments, margins"),
    (5, "StockAnalysis quarterly financials", urls["quarterly"], "Quarterly revenue, segment revenue, EPS"),
    (6, "StockAnalysis balance sheet (quarterly)", urls["balance"], "Cash, debt, unearned revenue, equity"),
    (7, "StockAnalysis cash flow", urls["cashflow"], "OCF, capex, issuance, SBC"),
    (8, "Nebius Q2 2026 shareholder letter (SEC 6-K Ex. 99.2)", urls["q2letter"], "ARR, adj. EBITDA, capex, deals, financing"),
    (9, "Nebius Q2 2026 press release", urls["q2pr"], "Headline results"),
    (10, "TradingKey Q2 2026 call summary", urls["q2call"], "Reaffirmed FY2026 guidance"),
    (11, "Nebius Q1 2026 earnings call transcript", urls["q1call"], "Capex guide raised to $20-$25B"),
    (12, "Nebius Q4 2025 shareholder letter", urls["q4letter"], "ARR $7-$9B target"),
    (13, "Nebius 6-K: Citi bridge facility", urls["citi"], "Facility terms"),
    (14, "Nebius press release: Citi facility", urls["citipr"], "Maturity, pricing, $40B commitments"),
    (15, "Nebius 6-K: Microsoft agreement", urls["msft"], "Contract value"),
    (16, "Nebius convertible notes pricing", urls["converts"], "2030/2032 notes"),
    (17, "CNBC: Meta $27B agreement", urls["meta"], "Contract value"),
    (18, "Yahoo Finance: CEO on tighter neocloud financing", urls["ceo"], "October drawdown, financing climate"),
    (19, "CNBC US 10-Year Treasury", urls["treasury"], "WACC risk-free rate"),
]
for r, row in enumerate(source_rows, 4):
    for c, value in enumerate(row, 1):
        put(ws, r, c, value, wrap=True)
ws.column_dimensions["A"].width = 7
ws.column_dimensions["B"].width = 44
ws.column_dimensions["C"].width = 95
ws.column_dimensions["D"].width = 45

for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = False

wb.save(OUT)
check = load_workbook(OUT, data_only=False)
expected = ["Valuation", "WACC", "Scenarios", "Actuals Source Audit", "Questions", "Sources"]
assert check.sheetnames == expected, check.sheetnames
assert abs(check["Scenarios"]["C19"].value - weighted_fv) < 1e-6
assert scen["Bear"]["target"] < PRICE < scen["Base"]["target"] < scen["Bull"]["target"]
print(f"WACC: {wacc:.2%} (Ke {ke:.2%}, Kd {cost_debt:.2%})")
print(f"EV ${PROVIDER_EV:,.0f}M; net debt ${NET_DEBT:,.0f}M; EV/FY27E sales {PROVIDER_EV / REV_FY27_CONS:.2f}x")
for k, s in scen.items():
    print(f"{k}: rev CAGR {s['cagr']:.1%}, EBITDA ${s['ebitda']:,.0f}M, EV ${s['ev']:,.0f}M, target ${s['target']:.2f} ({s['target'] / PRICE - 1:.1%}), PV ${s['pv']:.2f}")
print(f"Probability-weighted FV: ${weighted_fv:.2f} ({weighted_fv / PRICE - 1:.1%}); PV ${weighted_pv:.2f} ({weighted_pv / PRICE - 1:.1%})")
print(f"Analyst target cross-check: ${ANALYST_PT}")
print(f"Created {OUT}")
