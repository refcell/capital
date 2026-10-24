#!/usr/bin/env python3
"""Build the 2026-10-24 Carlyle Group valuation workbook."""

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


OUT = Path(__file__).with_name("2026-10-24 Carlyle Group Model.xlsx")
PRICE = 38.59  # Oct. 23, 2026 close
SHARES = 356.33  # millions
MARKET_CAP = 13750.0  # $ millions
PROVIDER_EV = 25960.0  # includes consolidated-fund CLO loans
CORP_CASH = 1256.5  # Jun. 30, 2026, excludes consolidated funds
CORP_DEBT = 2648.5  # excludes ~$350M CLO loans and $451M leases
CORP_NET_DEBT = CORP_DEBT - CORP_CASH
CORP_EV = MARKET_CAP + CORP_NET_DEBT
CONSOL_FUND_LOANS = 10804.4
INVESTMENTS = 3352.4  # attributable to CG shareholders, net of CLO loans
NET_ACCRUED_CARRY = 2414.6
FRE_TTM = 311.9 + 290.4 + 300.0 + 357.7
DE_PRETAX_TTM = 368.4 + 436.4 + 327.0 + 472.3
DE_PS_TTM = 0.96 + 1.01 + 0.89 + 1.07
DE_PS_FY25 = 4.02
EPS_FY26_CONS = 3.54
EPS_FY27_CONS = 4.98
GAAP_EPS_TTM = 0.97
DIVIDEND = 1.40
TAX = 0.21
FRE_AT_PS = FRE_TTM * (1 - TAX) / SHARES
CARRY_AT_PS = NET_ACCRUED_CARRY * (1 - TAX) / SHARES
INVEST_PS = INVESTMENTS / SHARES
NET_DEBT_PS = CORP_NET_DEBT / SHARES
SOTP_FRE_MULT = 15.0
SOTP = FRE_AT_PS * SOTP_FRE_MULT + CARRY_AT_PS + INVEST_PS - NET_DEBT_PS
ANALYST_PT = 52.21

NAVY = "17365D"
BLUE = "D9EAF7"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
WHITE = "FFFFFF"
thin = Side(style="thin", color="B7B7B7")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def put(ws, row, col, value, *, bold=False, fill=None, fmt=None, wrap=False):
    cell = ws.cell(row=row, column=col, value=value)
    cell.border = border
    cell.font = Font(bold=bold, color=WHITE if fill == NAVY else "000000")
    cell.alignment = Alignment(vertical="top", wrap_text=wrap)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if fmt:
        cell.number_format = fmt
    return cell


def title(ws, text, end_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    cell = ws.cell(1, 1, text)
    cell.font = Font(size=15, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 24


def headers(ws, row, values):
    for col, value in enumerate(values, 1):
        put(ws, row, col, value, bold=True, fill=BLUE, wrap=True)


# Scenario inputs: five-year horizon to FY2031, after-tax distributable earnings (DE) per share x exit P/E.
FRE_FY25 = 1236.2
scen = {
    "Bear": {"fre_cagr": 0.01, "de_ps": 3.60, "pe": 9.0, "fre_mult": 11.0, "carry_ps": 3.00, "shares": 350.0, "divs": 7.00, "w": 0.30},
    "Base": {"fre_cagr": 0.08, "de_ps": 5.40, "pe": 11.5, "fre_mult": 15.0, "carry_ps": 6.00, "shares": 335.0, "divs": 7.00, "w": 0.50},
    "Bull": {"fre_cagr": 0.12, "de_ps": 7.00, "pe": 13.0, "fre_mult": 19.0, "carry_ps": 9.00, "shares": 320.0, "divs": 7.50, "w": 0.20},
}
for s in scen.values():
    s["fre"] = FRE_FY25 * (1 + s["fre_cagr"]) ** 6  # FY2025 base to FY2031
    s["target"] = s["de_ps"] * s["pe"]
    s["de_cagr"] = (s["de_ps"] / DE_PS_TTM) ** (1 / 5) - 1
    s["fre_ps_at"] = s["fre"] * (1 - TAX) / s["shares"]
    s["sotp"] = s["fre_ps_at"] * s["fre_mult"] + s["carry_ps"] + INVEST_PS - NET_DEBT_PS
weighted_fv = sum(s["target"] * s["w"] for s in scen.values())

wb = Workbook()
ws = wb.active
ws.title = "Valuation"
title(ws, "The Carlyle Group Inc. (NASDAQ: CG) — Valuation Summary", 4)
summary = [
    ("Company", "The Carlyle Group Inc.", "Model date", "2026-10-24"),
    ("Quote", "$38.59 at Oct. 23, 2026 close", "Shares outstanding", "356.33M"),
    ("Market capitalization", "$13.75B", "Corporate EV (ex consolidated funds)", f"${CORP_EV / 1000:.2f}B"),
    ("Corporate net debt", f"${CORP_NET_DEBT / 1000:.2f}B (Jun. 30, 2026)", "Provider EV", "$25.96B — includes $10.8B consolidated CLO loans; not used"),
    ("Primary lens", "P/E on after-tax distributable earnings (DE); FRE sum-of-the-parts cross-check", "Stance", "Watch / Positive"),
]
for r, row in enumerate(summary, 3):
    for c, value in enumerate(row, 1):
        put(ws, r, c, value, bold=c in (1, 3), fill=YELLOW if c in (1, 3) else None, wrap=True)

headers(ws, 10, ["Metric", "Current / Forward", "Interpretation", "Source date"])
metrics = [
    ("Trailing P/E (GAAP)", f"{PRICE / GAAP_EPS_TTM:.2f}x", "$0.97 TTM GAAP EPS; depressed by $681M Q1 2026 unrealized carry reversal", "2026-06-30"),
    ("P / TTM after-tax DE", f"{PRICE / DE_PS_TTM:.2f}x", f"${DE_PS_TTM:.2f} DE/share over Q3 2025-Q2 2026 (company non-GAAP)", "2026-06-30"),
    ("Forward P/E (provider NTM)", "10.03x", "Implies ~$3.85 NTM EPS on S&P normalized basis", "2026-10-23"),
    ("FY2026E P/E (consensus)", f"{PRICE / EPS_FY26_CONS:.2f}x", "$3.54 (8 analysts, $3.40-$3.77); implies only $1.58 in H2 vs $1.96 YTD DE", "2026-10-23"),
    ("FY2027E P/E (consensus)", f"{PRICE / EPS_FY27_CONS:.2f}x", "$4.98, +40.6%; requires realizations and fundraising super cycle", "2026-10-23"),
    ("Price / after-tax FRE per share", f"{PRICE / FRE_AT_PS:.1f}x", f"TTM FRE ${FRE_TTM:,.0f}M taxed at 21% = ${FRE_AT_PS:.2f}/share", "2026-06-30"),
    ("Corporate EV / pre-tax FRE", f"{CORP_EV / FRE_TTM:.1f}x", "Large-cap alt managers trade at materially higher FRE multiples", "2026-06-30"),
    ("P/B", "2.60x", "BVPS $14.83; TBVPS $13.59", "2026-06-30"),
    ("Dividend yield", f"{DIVIDEND / PRICE:.2%}", f"$1.40 fixed policy; {DIVIDEND / DE_PS_TTM:.0%} of TTM DE", "2026-10-23"),
    ("Net accrued carry / share (pre-tax)", f"${NET_ACCRUED_CARRY / SHARES:.2f}", "$2.41B; down from $2.86B at Dec. 31, 2025", "2026-06-30"),
    ("Balance sheet investments / share", f"${INVEST_PS:.2f}", "$3.35B attributable, net of CLO loans", "2026-06-30"),
    ("Corporate net debt / share", f"${NET_DEBT_PS:.2f}", "$2.65B debt less $1.26B cash; $1.0B revolver undrawn", "2026-06-30"),
    ("SOTP cross-check", f"${SOTP:.2f}", f"{SOTP_FRE_MULT:.0f}x after-tax FRE + after-tax carry + investments − net debt", "2026-10-24"),
    ("Short interest", "6.18% of shares", "22.02M shares; 6.6 days to cover", "2026-10-23"),
    ("Probability-weighted FV (5Y)", f"${weighted_fv:.2f}", f"30% bear / 50% base / 20% bull; {weighted_fv / PRICE - 1:.1%} price upside before dividends", "2026-10-24"),
]
for r, row in enumerate(metrics, 11):
    for c, value in enumerate(row, 1):
        put(ws, r, c, value, bold=r == 25, fill=GREEN if r == 25 else None, wrap=True)
ws.column_dimensions["A"].width = 32
ws.column_dimensions["B"].width = 30
ws.column_dimensions["C"].width = 72
ws.column_dimensions["D"].width = 16
ws.freeze_panes = "A11"

ws = wb.create_sheet("WACC")
title(ws, "CG Weighted Average Cost of Capital", 4)
headers(ws, 3, ["Component", "Value", "Calculation / Source", "Comment"])
rf = 0.04999
erp = 0.05
beta = 1.87
ke = rf + beta * erp
cost_debt = (76.7 * 2) / CORP_DEBT
equity_weight = MARKET_CAP / (MARKET_CAP + CORP_DEBT)
debt_weight = 1 - equity_weight
wacc = equity_weight * ke + debt_weight * cost_debt * (1 - TAX)
wacc_rows = [
    ("Risk-free rate", rf, "CNBC US10Y", "Oct. 23, 2026"),
    ("Equity risk premium", erp, "Model assumption", "Standard US ERP"),
    ("Levered beta", beta, "StockAnalysis statistics", "Five-year beta; reflects carry and market-sentiment leverage"),
    ("Cost of equity (CAPM)", ke, "Rf + beta × ERP", "Implied DE yield at $38.59 is ~10.2%"),
    ("Pre-tax cost of debt", cost_debt, "H1 2026 segment interest $76.7M × 2 / corporate debt", "Senior notes; consolidated CLO loans excluded"),
    ("Tax rate", TAX, "US statutory", "TTM GAAP effective rate 15.8%"),
    ("Market capitalization", MARKET_CAP, "StockAnalysis", "$ millions"),
    ("Corporate debt", CORP_DEBT, "Q2 2026 earnings presentation", "$ millions; excludes $10.8B consolidated-fund loans"),
    ("Equity weight", equity_weight, "MC / (MC + debt)", ""),
    ("Debt weight", debt_weight, "Debt / (MC + debt)", ""),
    ("Computed WACC", wacc, "E/V × Ke + D/V × Kd × (1 − t)", "Provider estimate 7.15%; high beta drives the gap"),
]
for r, row in enumerate(wacc_rows, 4):
    for c, value in enumerate(row, 1):
        fmt = None
        if c == 2:
            fmt = "#,##0.0" if r in (10, 11) else ("0.00" if r == 6 else "0.00%")
        put(ws, r, c, value, bold=r == 14, fill=GREEN if r == 14 else None, fmt=fmt, wrap=True)
ws.column_dimensions["A"].width = 28
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 48
ws.column_dimensions["D"].width = 60

ws = wb.create_sheet("Scenarios")
title(ws, "CG Five-Year Distributable Earnings Scenario Analysis", 6)
ws.merge_cells("A2:F2")
ws["A2"] = (
    "Primary framework: FY2031 after-tax DE per share × exit P/E. GAAP FCF is not used: consolidated CLO "
    "purchases drove -$4.1B TTM FCF and are not Carlyle's cash. The SOTP rows value FRE, carry, investments and "
    "corporate net debt separately as a cross-check."
)
ws["A2"].alignment = Alignment(wrap_text=True)
ws.row_dimensions[2].height = 45
headers(ws, 4, ["Metric", "Bear", "Base", "Bull", "Units / Formula", "Interpretation"])
B, M, U = scen["Bear"], scen["Base"], scen["Bull"]
scenario_rows = [
    ("FRE CAGR (FY2025-FY2031)", B["fre_cagr"], M["fre_cagr"], U["fre_cagr"], "%", "Bear: fee step-downs offset wealth growth; base: fundraising super cycle lifts fees and margin; bull: plus wealth and capital markets"),
    ("Terminal FRE (FY2031)", B["fre"], M["fre"], U["fre"], "$mm", "From $1,236M FY2025 FRE"),
    ("Terminal after-tax DE / share", B["de_ps"], M["de_ps"], U["de_ps"], "$ / share", f"TTM ${DE_PS_TTM:.2f}; FY2025 $4.02"),
    ("Implied DE/share CAGR (5Y)", B["de_cagr"], M["de_cagr"], U["de_cagr"], "%", "From TTM base"),
    ("Exit P/E on DE", B["pe"], M["pe"], U["pe"], "x", "Current ~9.8x TTM DE; 10.0x provider NTM"),
    ("Target price", B["target"], M["target"], U["target"], "$ / share", "DE/share × P/E"),
    ("Upside / (downside)", B["target"] / PRICE - 1, M["target"] / PRICE - 1, U["target"] / PRICE - 1, "%", "Versus $38.59 close, before dividends"),
    ("Cumulative dividends (5Y)", B["divs"], M["divs"], U["divs"], "$ / share", "$1.40 fixed policy; bull assumes an increase late in period"),
    ("Probability", B["w"], M["w"], U["w"], "%", "Weights sum to 100%"),
    ("Weighted value / share", B["target"] * B["w"], M["target"] * M["w"], U["target"] * U["w"], "$ / share", "Target × probability"),
    ("Probability-weighted FV", None, weighted_fv, None, "$ / share", "Sum of weighted values"),
    ("Upside from current", None, weighted_fv / PRICE - 1, None, "%", "Excludes dividends"),
    ("Cross-check: FY2031 shares", B["shares"], M["shares"], U["shares"], "mm", "$1.6B buyback capacity remaining; SBC dilution offsets"),
    ("Cross-check: after-tax FRE / share", B["fre_ps_at"], M["fre_ps_at"], U["fre_ps_at"], "$ / share", "Terminal FRE × (1 − 21%) / shares"),
    ("Cross-check: FRE multiple", B["fre_mult"], M["fre_mult"], U["fre_mult"], "x", f"Current ~{PRICE / FRE_AT_PS:.1f}x after-tax FRE (whole share price)"),
    ("Cross-check: carry + investments − net debt", B["carry_ps"] + INVEST_PS - NET_DEBT_PS, M["carry_ps"] + INVEST_PS - NET_DEBT_PS, U["carry_ps"] + INVEST_PS - NET_DEBT_PS, "$ / share", f"After-tax carry per case + ${INVEST_PS:.2f} investments − ${NET_DEBT_PS:.2f} net debt"),
    ("Cross-check: SOTP price", B["sotp"], M["sotp"], U["sotp"], "$ / share", "Bear SOTP sits near today's price: the market is pricing FRE at a deep discount"),
]
for r, row in enumerate(scenario_rows, 5):
    for c, value in enumerate(row, 1):
        fmt = None
        if c in (2, 3, 4) and row[4] == "%":
            fmt = "0.0%"
        elif c in (2, 3, 4) and row[4] == "$ / share":
            fmt = "$#,##0.00"
        elif c in (2, 3, 4) and row[4] in ("$mm", "mm"):
            fmt = "#,##0"
        elif c in (2, 3, 4) and row[4] == "x":
            fmt = "0.0x"
        put(ws, r, c, value, bold=r in (10, 15, 16), fill=GREEN if r in (15, 16) else None, fmt=fmt, wrap=True)
ws.column_dimensions["A"].width = 38
for col in "BCD":
    ws.column_dimensions[col].width = 14
ws.column_dimensions["E"].width = 16
ws.column_dimensions["F"].width = 72
ws.freeze_panes = "A5"

urls = {
    "overview": "https://stockanalysis.com/stocks/cg/",
    "income": "https://stockanalysis.com/stocks/cg/financials/",
    "quarterly": "https://stockanalysis.com/stocks/cg/financials/?p=quarterly",
    "balance": "https://stockanalysis.com/stocks/cg/financials/balance-sheet/",
    "cashflow": "https://stockanalysis.com/stocks/cg/financials/cash-flow-statement/",
    "statistics": "https://stockanalysis.com/stocks/cg/statistics/",
    "forecast": "https://stockanalysis.com/stocks/cg/forecast/",
    "q2pr": "https://www.sec.gov/Archives/edgar/data/1527166/000152716626000036/cg202606308-kexx992edgar.htm",
    "q2deck": "https://ir.carlyle.com/static-files/16dcc049-4625-4913-935a-8db22517f995",
    "q2call": "https://ir.carlyle.com/static-files/8744b2ed-82a0-4a61-8850-f3ca66f7505f",
    "q2news": "https://ir.carlyle.com/news-releases/news-release-details/carlyle-reports-second-quarter-2026-financial-results",
    "investing": "https://www.investing.com/news/transcripts/earnings-call-transcript-carlyle-group-tops-q2-2026-estimates-as-stock-rises-premarket-93CH-4837839",
    "treasury": "https://www.cnbc.com/quotes/US10Y",
}
ws = wb.create_sheet("Actuals Source Audit")
title(ws, "CG Actuals Source Audit", 5)
headers(ws, 3, ["Data point", "Value", "Source URL", "As of", "Notes"])
audit = [
    ("Close price", "$38.59 (52-wk $35.41-$67.30)", urls["overview"], "2026-10-23", "After-hours $38.86; $50.64 close before Q2 report"),
    ("Market cap / provider EV", "$13.75B / $25.96B", urls["statistics"], "2026-10-23", "Provider EV includes consolidated-fund CLO loans"),
    ("Shares outstanding", "356.33M; +1.18% YoY", urls["statistics"], "2026-10-23", "Insiders 27.37%; float 258.61M"),
    ("Beta", "1.87", urls["statistics"], "2026-10-23", "Five-year beta"),
    ("GAAP revenue TTM / FY2025 / FY2024", "$2,797M / $4,032M / $4,740M", urls["income"], "2026-06-30", "Swings with unrealized performance allocations"),
    ("GAAP operating income / net income TTM", "$615.5M / $363.9M", urls["income"], "2026-06-30", "FY2025 $1,041M / $808.7M"),
    ("GAAP EPS TTM / FY2025 / FY2023", "$0.97 / $2.18 / -$1.68", urls["income"], "2026-06-30", "Q1 2026 -$0.37 on $681M carry reversal"),
    ("Q2 2026 GAAP", "Revenue $1,123.5M; pre-tax $273.8M; EPS $0.37", urls["q2pr"], "2026-06-30", "Effective tax rate 17.9%"),
    ("Q2 2026 DE / FRE", "DE $472.3M pre-tax, $1.07/share; FRE $357.7M (47%)", urls["q2pr"], "2026-06-30", "Best DE quarter in nearly four years"),
    ("Quarterly DE per share", "0.91 / 0.96 / 1.01 / 0.89 / 1.07", urls["q2deck"], "2Q25-2Q26", f"TTM ${DE_PS_TTM:.2f}; FY2025 $4.02"),
    ("Quarterly FRE ($M)", "323.3 / 311.9 / 290.4 / 300.0 / 357.7", urls["q2deck"], "2Q25-2Q26", f"TTM ${FRE_TTM:,.1f}M; FY2025 $1,236.2M; FY2024 $1,104.6M"),
    ("Segment fund management fees ($M)", "589.6 / 573.9 / 554.1 / 544.5 / 560.1", urls["q2deck"], "2Q25-2Q26", "Down 5% YoY in Q2 2026"),
    ("Transaction fees / FRPR ($M)", "Q2 2026: $110.5M / $88.7M", urls["q2deck"], "2026-06-30", "Both records; drove FRE beat"),
    ("Realized net performance revenues", "Q2 $114.6M; FY2025 $357.3M; FY2024 $366.1M", urls["q2deck"], "2026-06-30", "Japan buyout IV first carry"),
    ("AUM / FEAUM", "$485.5B (+4%) / $334.4B (+3%)", urls["q2pr"], "2026-06-30", "GPE $162.7B, GC $211.1B, AlpInvest $111.7B"),
    ("Perpetual / pending FEAUM", "$113B (34%) / $28B (+57%)", urls["q2pr"], "2026-06-30", "Available capital $97B"),
    ("Inflows / deployment / realizations (LTM)", "$55.8B / $53.0B / $36.8B", urls["q2pr"], "2026-06-30", "Carry fund appreciation 7% LTM"),
    ("Net accrued performance revenues", "$2,414.6M", urls["q2deck"], "2026-06-30", "GPE $1,476M; GC $283M; AlpInvest $655M; $2,859M at 4Q25"),
    ("Corporate cash / debt", "$1,256.5M / $2,648.5M", urls["q2deck"], "2026-06-30", "Excludes consolidated funds, $350M CLO loans and $451M leases"),
    ("Investments attributable to CG", "$3,352.4M", urls["q2deck"], "2026-06-30", "Net of CLO loans; excludes NGP stake"),
    ("Consolidated cash / debt (provider)", "$2,041M / $14,254M", urls["balance"], "2026-06-30", "Includes $784M fund cash and $10,804M fund loans"),
    ("Equity / BVPS / TBVPS", "$5,242M common; $14.83 / $13.59", urls["balance"], "2026-06-30", "NCI $1,971M"),
    ("GAAP OCF / capex TTM", "-$3,998M / -$125.3M", urls["cashflow"], "TTM 2026-06-30", "Consolidated CLO investment purchases flow through OCF"),
    ("Buybacks / dividends TTM", "$915.3M / $504.5M", urls["cashflow"], "TTM 2026-06-30", "Q2: 6.7M shares for $304M; $1.6B authorization left"),
    ("SBC TTM", "$414.1M", urls["cashflow"], "TTM 2026-06-30", "Excluded from DE; ~15% of segment fee revenue"),
    ("Consensus FY2026 / FY2027", "Rev $3.65B / $4.75B; EPS $3.54 / $4.98", urls["forecast"], "2026-10-23", "S&P Global; FY2026 8 analysts ($3.40-$3.77); FY2028 gated"),
    ("Analyst targets", "$52.21 avg; $41-$73; Buy", urls["forecast"], "2026-10-23", "20 analysts: 6 SB, 4 B, 9 H, 1 S; eight cuts in October"),
    ("Short interest", "22.02M; 6.18% of shares", urls["statistics"], "2026-10-23", "Up from 19.26M prior month"),
    ("Dividend", "$0.35 quarterly; $1.40 annual", urls["q2news"], "2026-08-05", "Fixed policy, not a payout ratio"),
    ("2026 comp ratio guide", "~47%", urls["q2call"], "2026-08-05", "Margin expansion expected 2027-2028"),
    ("Next earnings", "Nov. 5, 2026", urls["overview"], "2026-10-23", "Q3 2026"),
    ("US 10Y Treasury", "4.999%", urls["treasury"], "2026-10-23", "CNBC quote page"),
]
for r, row in enumerate(audit, 4):
    for c, value in enumerate(row, 1):
        put(ws, r, c, value, wrap=True)
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 42
ws.column_dimensions["C"].width = 70
ws.column_dimensions["D"].width = 16
ws.column_dimensions["E"].width = 55
ws.freeze_panes = "A4"

ws = wb.create_sheet("Questions")
title(ws, "CG Open Diligence Questions", 4)
headers(ws, 3, ["#", "Question", "Why it matters", "Evidence needed"])
questions = [
    (1, "Why does FY2026 consensus EPS of $3.54 imply only ~$1.58 of H2 DE when YTD DE is $1.96 and H2 2025 was $1.97?", "Either the street expects a realization drought or the provider EPS basis differs from company DE.", "Q3 DE vs consensus; sell-side model definitions"),
    (2, "Why are segment fund management fees down 5% YoY while FEAUM is up 3%?", "Fee rate compression or GPE fund step-downs undercut the recurring base.", "Fee rate by segment; step-down schedule for CP VIII and older funds"),
    (3, "How much of the record $110.5M transaction fee quarter is repeatable?", "Capital markets fees are cyclical and lifted the Q2 FRE beat.", "Quarterly transaction fees through 2027"),
    (4, "When will the next U.S. buyout fund (after the $5B anchor) hold a first close, and at what size and fee rate?", "Flagship fee activation is the main FRE growth lever in 2027-2028.", "Fundraising disclosures; marketing size"),
    (5, "What drove the 15% drop in net accrued carry to $2.41B since December?", "Lower accrued carry in U.S. buyout VII reduces future realized DE.", "Fund-level marks and carry status"),
    (6, "How exposed are credit and PE portfolios to software and AI-disrupted borrowers?", "Alt-manager selloff (~27% YTD per Morgan Stanley) is partly a private-credit quality scare.", "Sector exposure, non-accruals in direct lending"),
    (7, "What are the economics and timing of the Unum reinsurance transaction?", "Analysts expect it to lift insurance-solutions AUM (~$86-87B, flat YTD) when it closes later in 2026.", "Closing 8-K, AUM and fee rate"),
    (8, "What was committed, spent or risked on the abandoned Lukoil international asset deal?", "Reputational and sanctions risk; possible broken-deal costs.", "Fund documents, cost disclosure"),
    (9, "What share of the $3.35B balance sheet investments is in CLO risk retention and seed capital, and what are their marks?", "Investments are ~$9.41/share of value in the SOTP.", "Investment breakdown in 10-Q"),
    (10, "Will SBC ($414M TTM) keep rising with the wealth build-out and AI investment?", "DE excludes equity comp; real per-share cost appears as dilution.", "Equity grant schedule, diluted share count"),
    (11, "Can the 47% FRE margin rise in 2027-2028 as management expects?", "Base case needs margin expansion; comp ratio guided flat for 2026.", "2027 FRE margin trajectory"),
    (12, "How fast is evergreen wealth AUM ($20B, +60%) growing and what are redemption terms?", "Wealth products add FRPR but can face gating during stress.", "Monthly flows, redemption queues"),
    (13, "Is the $1.40 dividend policy likely to be raised once DE grows?", "Fixed dividend at 36% of DE leaves capacity; affects total return.", "Board commentary"),
    (14, "How will consolidated CLO growth affect reported GAAP debt and OCF?", "Provider screens show $14.3B debt and -$4.1B FCF that are not corporate obligations.", "Consolidated vs operating-entity balance sheet each quarter"),
    (15, "What are the three-year strategic plan targets from Investor Day and is the company on track?", "Management cites the plan as its anchor; targets were not reviewed here.", "Investor Day materials"),
    (16, "What is Carlyle's dependence on Japan buyout and U.S. buyout VI/VII for 2026-2027 realized carry?", "Carry concentration in a few funds makes DE lumpy.", "Fund-level net accrued carry and realization pipeline"),
]
for r, row in enumerate(questions, 4):
    for c, value in enumerate(row, 1):
        put(ws, r, c, value, wrap=True)
ws.column_dimensions["A"].width = 7
ws.column_dimensions["B"].width = 72
ws.column_dimensions["C"].width = 60
ws.column_dimensions["D"].width = 45

ws = wb.create_sheet("Sources")
title(ws, "CG Sources", 4)
headers(ws, 3, ["#", "Source", "URL", "Use"])
source_rows = [
    (1, "StockAnalysis overview", urls["overview"], "Price, range, news, earnings date"),
    (2, "StockAnalysis statistics", urls["statistics"], "Market value, ratios, shares, beta, short interest"),
    (3, "StockAnalysis forecast", urls["forecast"], "Consensus revenue/EPS, targets and ratings"),
    (4, "StockAnalysis income statement", urls["income"], "GAAP annual revenue, segments, margins"),
    (5, "StockAnalysis quarterly financials", urls["quarterly"], "Quarterly GAAP revenue and EPS"),
    (6, "StockAnalysis balance sheet", urls["balance"], "Consolidated cash, debt, equity"),
    (7, "StockAnalysis cash flow", urls["cashflow"], "OCF, capex, buybacks, dividends, SBC"),
    (8, "Carlyle Q2 2026 8-K Ex. 99.2 (SEC)", urls["q2pr"], "DE, FRE, AUM, segment results"),
    (9, "Carlyle Q2 2026 earnings presentation", urls["q2deck"], "Corporate balance sheet, carry, buybacks"),
    (10, "Carlyle Q2 2026 earnings call transcript", urls["q2call"], "Comp ratio, fundraising, wealth, capital markets"),
    (11, "Carlyle Q2 2026 news release", urls["q2news"], "Dividend, headline results"),
    (12, "Investing.com Q2 2026 call summary", urls["investing"], "Pre-report price, consensus comparison"),
    (13, "CNBC US 10-Year Treasury", urls["treasury"], "WACC risk-free rate"),
]
for r, row in enumerate(source_rows, 4):
    for c, value in enumerate(row, 1):
        put(ws, r, c, value, wrap=True)
ws.column_dimensions["A"].width = 7
ws.column_dimensions["B"].width = 38
ws.column_dimensions["C"].width = 95
ws.column_dimensions["D"].width = 50

for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = False

wb.save(OUT)
check = load_workbook(OUT, data_only=False)
expected = ["Valuation", "WACC", "Scenarios", "Actuals Source Audit", "Questions", "Sources"]
assert check.sheetnames == expected, check.sheetnames
assert check["Scenarios"]["C15"].value == weighted_fv
assert scen["Bear"]["target"] < PRICE < scen["Base"]["target"] < scen["Bull"]["target"]
print(f"WACC: {wacc:.2%} (Ke {ke:.2%}, Kd {cost_debt:.2%})")
print(f"Corporate EV ${CORP_EV:,.0f}M; net debt ${CORP_NET_DEBT:,.0f}M; EV/FRE {CORP_EV / FRE_TTM:.1f}x")
print(f"After-tax FRE/share ${FRE_AT_PS:.2f}; carry ${CARRY_AT_PS:.2f}; investments ${INVEST_PS:.2f}; net debt ${NET_DEBT_PS:.2f}; SOTP ${SOTP:.2f}")
for k, s in scen.items():
    print(f"{k}: FRE ${s['fre']:.0f}M, DE/sh CAGR {s['de_cagr']:.1%}, target ${s['target']:.2f} ({s['target'] / PRICE - 1:.1%}), SOTP ${s['sotp']:.2f}")
print(f"Probability-weighted fair value: ${weighted_fv:.2f} ({weighted_fv / PRICE - 1:.1%})")
print(f"Analyst target cross-check: ${ANALYST_PT}; base vs target {(scen['Base']['target'] / ANALYST_PT - 1):.1%}")
print(f"Created {OUT}")
