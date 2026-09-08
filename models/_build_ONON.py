#!/usr/bin/env python3
"""Build On Holding's six-sheet valuation workbook."""
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).with_name("2026-09-08 On Holding Model.xlsx")
PRICE, SHARES = 27.99, 334.21  # USD, millions
MARKET_CAP, ENTERPRISE_VALUE = 9350.0, 8520.0  # USD millions
NET_CASH = MARKET_CAP - ENTERPRISE_VALUE
FY26_REVENUE_USD = 4350.0  # CHF3.51B consensus translated at ~1.24 USD/CHF

SCENARIOS = {
    "Bear": {"growth": .08, "margin": .10, "multiple": 15, "net_cash": 600, "shares": 350, "discount": .12, "weight": .25},
    "Base": {"growth": .14, "margin": .14, "multiple": 21, "net_cash": 850, "shares": 340, "discount": .105, "weight": .55},
    "Bull": {"growth": .20, "margin": .17, "multiple": 27, "net_cash": 1100, "shares": 335, "discount": .095, "weight": .20},
}
for case in SCENARIOS.values():
    case["terminal_revenue"] = FY26_REVENUE_USD * (1 + case["growth"]) ** 5
    case["terminal_fcf"] = case["terminal_revenue"] * case["margin"]
    case["ev"] = case["terminal_fcf"] * case["multiple"]
    case["terminal_price"] = (case["ev"] + case["net_cash"]) / case["shares"]
    case["target"] = case["terminal_price"] / (1 + case["discount"]) ** 5
    case["upside"] = case["target"] / PRICE - 1
weighted = sum(c["target"] * c["weight"] for c in SCENARIOS.values())
assert SCENARIOS["Bear"]["target"] < PRICE
assert abs(SCENARIOS["Base"]["target"] / 44.63 - 1) < .20

rf, erp, beta = .0478, .05, 2.08
ke = rf + beta * erp
debt, kd, tax = 696.16, .045, .15
ew = MARKET_CAP / (MARKET_CAP + debt)
dw = 1 - ew
wacc = ew * ke + dw * kd * (1 - tax)

wb = Workbook()
navy, blue, gold, white = "17365D", "D9EAF7", "D8B34B", "FFFFFF"
thin = Side(style="thin", color="A6A6A6")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def put(ws, r, c, value, bold=False, fill=None, color="000000", size=10):
    x = ws.cell(r, c, value)
    x.font = Font(name="Aptos", size=size, bold=bold, color=color)
    x.fill = PatternFill("solid", fgColor=fill) if fill else PatternFill(fill_type=None)
    x.border = border
    x.alignment = Alignment(vertical="top", wrap_text=True)
    return x

def title(ws, text, cols, subtitle):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    put(ws, 1, 1, text, True, navy, white, 15)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=cols)
    put(ws, 2, 1, subtitle, True, blue)

def header(ws, r, values):
    for c, v in enumerate(values, 1): put(ws, r, c, v, True, navy, white)

def widths(ws, values):
    for c, v in enumerate(values, 1): ws.column_dimensions[get_column_letter(c)].width = v
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False

# Valuation
ws = wb.active
ws.title = "Valuation"
title(ws, "On Holding AG (ONON) — Valuation", 4, "Quote: September 4, 2026 close | Model date: September 8, 2026 | USD unless noted")
header(ws, 3, ["Field", "Value", "Comment"])
facts = [
    ("Company", "On Holding AG", "Swiss premium performance footwear and apparel brand"),
    ("Ticker", "NYSE: ONON", "Financial statements are reported in CHF; market valuation in USD"),
    ("Price", PRICE, "September 4 close"), ("Shares outstanding (M)", SHARES, "Current total shares"),
    ("Market capitalization ($M)", MARKET_CAP, "StockAnalysis"), ("Enterprise value ($M)", ENTERPRISE_VALUE, "StockAnalysis"),
    ("Net cash ($M)", NET_CASH, "Market cap less enterprise value; rounded"),
    ("Primary lens", "Discounted terminal FCF", "Forward adjusted P/E is the main cross-check"),
    ("Stance", "Watch / selective buy", "Strong growth and brand economics; Q2 miss and guidance reset require confirmation"),
]
for r, row in enumerate(facts, 4):
    for c, v in enumerate(row, 1): put(ws, r, c, v, c == 1)
header(ws, 14, ["Valuation metric", "Value", "Interpretation"])
metrics = [
    ("Trailing P/E", 19.11, "GAAP TTM"), ("Forward P/E", 15.14, "Provider adjusted estimates"),
    ("P/S", 2.35, "USD market cap / USD-converted revenue"), ("P/FCF", 17.31, "5.78% FCF yield"),
    ("EV/FCF", 15.76, "Net-cash-adjusted"), ("EV/Sales", 2.14, "Premium brand valuation"),
    ("EV/EBITDA", 14.03, "Secondary cross-check"), ("Analyst average target", 44.63, "28 analysts; not intrinsic value"),
]
for r, row in enumerate(metrics, 15):
    for c, v in enumerate(row, 1): put(ws, r, c, v, c == 1)
widths(ws, [31, 25, 96, 14])

# WACC
ws = wb.create_sheet("WACC")
title(ws, "On Holding — Weighted Average Cost of Capital", 4, "CAPM using market-value weights")
header(ws, 3, ["Component", "Value", "Source / formula"])
rows = [("Risk-free rate", rf, "U.S. 10-year reference, September 2026"), ("Equity risk premium", erp, "Model assumption"),
        ("Levered beta", beta, "StockAnalysis five-year beta"), ("Cost of equity", ke, "Rf + beta × ERP"),
        ("Pre-tax cost of debt", kd, "Normalized assumption; debt is primarily leases"), ("Tax rate", tax, "Normalized; TTM 7.95% is unusually low"),
        ("Market cap ($M)", MARKET_CAP, "USD millions"), ("Debt ($M)", debt, "USD millions"),
        ("Equity weight", ew, "E/(D+E)"), ("Debt weight", dw, "D/(D+E)"), ("WACC", wacc, "E/V×Ke + D/V×Kd×(1−T)")]
for r, row in enumerate(rows, 4):
    for c, v in enumerate(row, 1):
        x = put(ws, r, c, v, c == 1 or row[0] == "WACC", gold if row[0] == "WACC" else None)
        if c == 2 and row[0] not in {"Levered beta", "Market cap ($M)", "Debt ($M)"}: x.number_format = "0.00%"
widths(ws, [31, 22, 96, 14])

# Scenarios
ws = wb.create_sheet("Scenarios")
title(ws, "On Holding — Discounted FCF Scenarios", 6, "Five years from FY2026 consensus; USD millions")
header(ws, 3, ["Metric", "Bear", "Base", "Bull", "Notes"])
order = ("Bear", "Base", "Bull")
rows = [
    ("FY2026 revenue anchor", FY26_REVENUE_USD, FY26_REVENUE_USD, FY26_REVENUE_USD, "CHF3.51B consensus translated at ~1.24 USD/CHF"),
    ("Revenue CAGR (5Y)", *(SCENARIOS[k]["growth"] for k in order), "Brand growth after FY2026"),
    ("Terminal revenue", *(SCENARIOS[k]["terminal_revenue"] for k in order), "FY2026 anchor compounded five years"),
    ("Adjusted FCF margin", *(SCENARIOS[k]["margin"] for k in order), "TTM is 13.6%"),
    ("Terminal FCF", *(SCENARIOS[k]["terminal_fcf"] for k in order), "Revenue × margin"),
    ("Exit FCF multiple", *(SCENARIOS[k]["multiple"] for k in order), "Compression / durable growth / leadership"),
    ("Implied EV", *(SCENARIOS[k]["ev"] for k in order), "Terminal FCF × multiple"),
    ("Net cash", *(SCENARIOS[k]["net_cash"] for k in order), "Future capital position"),
    ("Shares outstanding", *(SCENARIOS[k]["shares"] for k in order), "SBC and issuance assumptions"),
    ("Terminal price", *(SCENARIOS[k]["terminal_price"] for k in order), "(EV + cash) / shares"),
    ("Discount rate", *(SCENARIOS[k]["discount"] for k in order), "Scenario risk"),
    ("Present target price", *(SCENARIOS[k]["target"] for k in order), "Five-year terminal value discounted"),
    ("Upside / downside", *(SCENARIOS[k]["upside"] for k in order), "Versus current price"),
    ("Probability", *(SCENARIOS[k]["weight"] for k in order), "25% / 55% / 20%"),
    ("Weighted value/share", *(SCENARIOS[k]["target"] * SCENARIOS[k]["weight"] for k in order), "Contribution")]
for r, row in enumerate(rows, 4):
    for c, v in enumerate(row, 1):
        x = put(ws, r, c, v, c == 1)
        if c in (2,3,4) and row[0] in {"Revenue CAGR (5Y)", "Adjusted FCF margin", "Discount rate", "Upside / downside", "Probability"}: x.number_format = "0.0%"
        elif c in (2,3,4) and row[0] in {"Terminal price", "Present target price", "Weighted value/share"}: x.number_format = "$0.00"
put(ws, 20, 1, "Probability-weighted fair value", True, gold)
put(ws, 20, 2, weighted, True, gold).number_format = "$0.00"
put(ws, 21, 1, "Upside from current price", True, gold)
put(ws, 21, 2, weighted / PRICE - 1, True, gold).number_format = "0.0%"
widths(ws, [31, 18, 18, 18, 96, 14])

# Actuals Source Audit
ws = wb.create_sheet("Actuals Source Audit")
title(ws, "On Holding — Actuals Source Audit", 5, "Market figures USD; financial statements CHF unless noted")
header(ws, 3, ["Data point", "Value", "Source URL", "Source date", "Notes"])
sa = "https://stockanalysis.com/stocks/onon"
audit = [
    ("Price", "$27.99", f"{sa}/", "2026-09-04", "Close"), ("Market cap / EV", "$9.35B / $8.52B", f"{sa}/statistics/", "2026-09-07", "USD"),
    ("Shares", "334.21M", f"{sa}/statistics/", "2026-09-07", "+0.95% YoY"), ("TTM revenue", "CHF3.220B", f"{sa}/financials/", "2026-06-30", "+18.48%"),
    ("TTM gross profit", "CHF2.088B", f"{sa}/financials/", "2026-06-30", "64.82% margin"), ("TTM operating income", "CHF443.5M", f"{sa}/financials/", "2026-06-30", "13.77% margin"),
    ("TTM net income", "CHF396.2M", f"{sa}/financials/", "2026-06-30", "12.30% margin"), ("TTM OCF / FCF", "CHF525.4M / CHF437.9M", f"{sa}/financials/cash-flow-statement/", "2026-06-30", "FCF margin 13.60%"),
    ("TTM capex", "CHF87.5M", f"{sa}/financials/cash-flow-statement/", "2026-06-30", "Cash outflow"), ("TTM SBC", "CHF72.3M", f"{sa}/financials/cash-flow-statement/", "2026-06-30", "16.5% of FCF"),
    ("Cash / debt", "CHF1.239B / CHF562.5M", f"{sa}/financials/balance-sheet/", "2026-06-30", "Net cash CHF676.9M"), ("Inventory", "CHF472.9M", f"{sa}/financials/balance-sheet/", "2026-06-30", "Higher than FY2025"),
    ("FY2026 revenue consensus", "CHF3.51B", f"{sa}/forecast/", "2026-09-03", "+16.54%; 27 analysts"), ("FY2026 adjusted EPS", "CHF1.41", f"{sa}/forecast/", "2026-09-03", "Non-GAAP adjusted"),
    ("FY2027 headline revenue", "CHF4.17B", f"{sa}/forecast/", "2026-09-03", "+18.62%; detailed table gated"), ("FY2027 adjusted EPS", "CHF1.68", f"{sa}/forecast/", "2026-09-03", "+18.92%; headline"),
    ("Target range", "$19.82-$73.28", f"{sa}/forecast/", "2026-09-03", "Average $44.63"), ("Beta", "2.08", f"{sa}/statistics/", "2026-09-07", "Five-year"),
    ("Latest earnings", "2026-08-11", f"{sa}/", "2026-08-11", "Q2 already released; next quarter is future catalyst")]
for r, row in enumerate(audit, 4):
    for c, v in enumerate(row, 1): put(ws, r, c, v, c == 1)
widths(ws, [35, 26, 100, 16, 80])

# Questions
ws = wb.create_sheet("Questions")
title(ws, "On Holding — Open Questions", 4, "Items that can change conviction or valuation")
header(ws, 3, ["#", "Question", "Why it matters", "Best next evidence"])
questions = [
    (1, "Was the Q2 sales miss demand weakness, deliberate wholesale restraint, or shipment timing?", "The answer determines whether the guide reset is temporary or structural.", "Q3 channel growth, orders and management bridge."),
    (2, "Can DTC growth above wholesale growth continue without higher acquisition cost?", "DTC supports margin and customer data but carries store and marketing costs.", "DTC cohort economics and contribution margin."),
    (3, "How much U.S. tariff cost is embedded in the revised outlook?", "Tariffs can pressure price, demand, or gross margin.", "Price/cost bridge and sourcing actions."),
    (4, "Can gross margin remain at least 65%?", "Premium economics and valuation depend on mix and full-price sell-through.", "Quarterly gross margin and discounting."),
    (5, "Why did inventory rise to CHF472.9M while guidance moderated?", "Excess inventory can force promotions and working-capital use.", "Inventory growth versus sales and aged stock."),
    (6, "What portion of 2026 growth comes from footwear versus apparel?", "Apparel is strategic optionality but may carry different economics.", "Category revenue and gross margin."),
    (7, "How concentrated is manufacturing by country and supplier?", "Tariffs and disruption can impair availability and margin.", "Supplier and origin disclosure."),
    (8, "Will the Federer association remain economically important after product expansion?", "Brand concentration around ambassadors can be both moat and risk.", "Campaign efficiency and product sell-through."),
    (9, "Why has share count continued to rise despite positive FCF?", "Dilution reduces per-share compounding.", "SBC, option exercises and repurchase policy."),
    (10, "Is the TTM 7.95% tax rate repeatable?", "A normalized tax rate lowers sustainable EPS and FCF.", "Jurisdictional tax bridge."),
    (11, "How much of lease debt relates to owned retail expansion?", "Retail stores improve DTC but add fixed costs and lease obligations.", "Lease maturity and store productivity."),
    (12, "What customer or wholesale partner concentration exists?", "Retailer resets could create abrupt shipment volatility.", "Top-ten customer share and order trends."),
    (13, "Can premium pricing hold during cautious consumer spending?", "Full-price demand is central to gross margin and brand equity.", "ASP, promotions and sell-through."),
    (14, "What is management's next earnings date and explicit FY2026 bridge?", "The next report must validate revised low-20s constant-currency growth.", "Investor relations calendar and Q3 release.")]
for r, row in enumerate(questions, 4):
    for c, v in enumerate(row, 1): put(ws, r, c, v, c == 1)
widths(ws, [7, 84, 74, 72])

# Sources
ws = wb.create_sheet("Sources")
title(ws, "On Holding — Sources", 4, "Public sources accessed September 8, 2026")
header(ws, 3, ["#", "Source", "URL", "Use"])
sources = [(1, "StockAnalysis overview", f"{sa}/", "Identity, price, profile, news and earnings"),
           (2, "StockAnalysis financials", f"{sa}/financials/", "Historical financials, segments and margins"),
           (3, "StockAnalysis balance sheet", f"{sa}/financials/balance-sheet/", "Cash, debt, inventory and equity"),
           (4, "StockAnalysis cash flow", f"{sa}/financials/cash-flow-statement/", "OCF, capex, FCF, SBC and working capital"),
           (5, "StockAnalysis statistics", f"{sa}/statistics/", "Valuation, shares, beta and returns"),
           (6, "StockAnalysis forecast", f"{sa}/forecast/", "Consensus, ratings and targets"),
           (7, "On Q2 FY2026 results", "https://investors.on-running.com/financials-and-filings/quarterly-results", "Official earnings materials and guidance"),
           (8, "U.S. Treasury yield curve", "https://home.treasury.gov/resource-center/data-chart-center/interest-rates", "Risk-free-rate reference")]
for r, row in enumerate(sources, 4):
    for c, v in enumerate(row, 1): put(ws, r, c, v, c == 1)
widths(ws, [7, 40, 112, 74])

for ws in wb.worksheets:
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, (int, float)) and cell.number_format == "General": cell.number_format = "#,##0.00"
wb.save(OUT)

check = load_workbook(OUT, data_only=False)
assert check.sheetnames == ["Valuation", "WACC", "Scenarios", "Actuals Source Audit", "Questions", "Sources"]
assert all(check[s].max_row >= 10 for s in check.sheetnames)
print(f"WACC: {wacc:.2%}")
print("Targets:", ", ".join(f"{k} ${v['target']:.2f}" for k, v in SCENARIOS.items()))
print(f"Weighted FV: ${weighted:.2f} ({weighted / PRICE - 1:+.1%})")
print(f"Wrote {OUT}")
